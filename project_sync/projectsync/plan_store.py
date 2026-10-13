"""計畫表（Tasks / Groups / People / Holidays / Settings / _sync_state）的讀寫。

三種來源：
- xlsx：本機檔案，可讀寫（無 Google 憑證時的主要模式）。
- google：Google Sheets（需 service account JSON），可讀寫。
- export：Google 公開匯出 `/export?format=xlsx`（「知道連結者可檢視」即可），唯讀。
"""

from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

import requests

from . import plan_core as pc

TASKS = "Tasks"
GROUPS = "Groups"
PEOPLE = "People"
HOLIDAYS = "Holidays"
SETTINGS = "Settings"
README = "說明"
GANTT = "Gantt"
STATE = "_sync_state"

TASK_HEADERS = [
    "ID", "WBS", "Group", "Task Name", "Type", "Priority", "Owner", "Duration (d)", "Predecessors",
    "Start Constraint", "Start", "Finish", "Status", "% Complete", "Sync to Jira", "Jira Key",
    "Description / 工作內容", "Discussion / Comments", "Proposed By", "Estimate", "Notes",
    "Sync Conflict", "Last Sync",
]
DATE_HEADERS = {"Start Constraint", "Start", "Finish"}
NUMERIC_HEADERS = {"Duration (d)", "% Complete"}

# 欄位別名 → 標準欄名（使用者改欄名/調順序也能讀）
_HEADER_ALIASES = {
    "id": "ID", "wbs": "WBS", "group": "Group", "group/epic": "Group", "epic": "Group",
    "task name": "Task Name", "task": "Task Name", "name": "Task Name", "type": "Type",
    "priority": "Priority", "owner": "Owner", "duration (d)": "Duration (d)", "duration": "Duration (d)",
    "predecessors": "Predecessors", "start constraint": "Start Constraint", "start": "Start",
    "finish": "Finish", "status": "Status", "% complete": "% Complete", "sync to jira": "Sync to Jira",
    "sync": "Sync to Jira", "jira key": "Jira Key", "description / 工作內容": "Description / 工作內容",
    "description": "Description / 工作內容", "工作內容": "Description / 工作內容",
    "discussion / comments": "Discussion / Comments", "discussion": "Discussion / Comments",
    "comments": "Discussion / Comments", "proposed by": "Proposed By", "estimate": "Estimate",
    "notes": "Notes", "sync conflict": "Sync Conflict", "last sync": "Last Sync",
}

HEADER_NOTES = {
    "ID": "唯一且固定不變，格式 PLN-<群組>-<序號>（例 PLN-AT-01）。新增列可留空，同步時自動編號。",
    "Group": "群組代碼，對應 Groups 分頁（每組 = 一個 Jira Epic）。",
    "Type": "Task / Milestone / Procurement。Duration=0 一律視為 Milestone。",
    "Priority": "P1 / P2 / P3；同步到 Jira 標籤 prio-P1 等。",
    "Owner": "對應 People 分頁的 Name；空白 = 預設受託人。",
    "Duration (d)": "工作天（Mon–Fri，扣除 Holidays）。0 = 里程碑。",
    "Predecessors": "MS Project 語法：PLN-AT-01FS+1d; PLN-FX-02SS; PLN-WZ-03FF-2d。類型省略=FS；lag 可用 d 或 w。",
    "Start Constraint": "選填。日期 = 不早於此日開始 (SNET)；'MSO 2026-11-02' = 必須於此日開始。",
    "Start": "程式計算，請勿手動修改（會被覆寫）。",
    "Finish": "程式計算，請勿手動修改（會被覆寫）。",
    "Sync to Jira": "Y = 已核准，同步到 Jira；N 或空白 = 草稿/討論中，只出現在 Confluence。",
    "Jira Key": "程式回寫，請勿修改。",
    "Description / 工作內容": "詳細工作範圍，同步到 Jira 描述。",
    "Discussion / Comments": "討論用，不會送到 Jira；會顯示在 Confluence 計畫頁。",
    "Estimate": "Y = 估計值，待確認。",
    "Sync Conflict": "兩邊都改了同一欄位時程式寫在這裡。輸入 use sheet 或 use jira 決定以哪邊為準。",
}


def canonical_header(h: Any) -> str:
    key = re.sub(r"\s+", " ", str(h or "").strip()).casefold()
    return _HEADER_ALIASES.get(key, str(h or "").strip())


def cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


@dataclass
class PlanBook:
    headers: list[str]
    rows: list[list[str]]  # data rows; sheet row number = index + 2
    groups: list[dict[str, str]] = field(default_factory=list)
    people: list[dict[str, str]] = field(default_factory=list)
    holidays: list[date] = field(default_factory=list)
    settings: dict[str, str] = field(default_factory=dict)
    state: dict[str, Any] = field(default_factory=dict)
    writable: bool = False
    source_desc: str = ""

    def col(self, header: str) -> int:
        try:
            return self.headers.index(header)
        except ValueError:
            return -1

    def row_dicts(self) -> list[tuple[int, dict[str, str]]]:
        out = []
        for i, r in enumerate(self.rows):
            d = {h: (r[j] if j < len(r) else "") for j, h in enumerate(self.headers)}
            out.append((i + 2, d))
        return out


@dataclass
class SheetChanges:
    cells: list[tuple[int, str, Any]] = field(default_factory=list)  # (row number, header, value)
    appends: list[dict[str, Any]] = field(default_factory=list)
    settings: dict[str, str] = field(default_factory=dict)
    state: dict[str, Any] | None = None

    def is_empty(self) -> bool:
        return not (self.cells or self.appends or self.settings or self.state is not None)


# ---------------------------------------------------------------------------
# Parse helpers shared by all stores
# ---------------------------------------------------------------------------


def _table(values: list[list[Any]]) -> tuple[list[str], list[list[str]]]:
    if not values:
        return [], []
    headers = [canonical_header(h) for h in values[0]]
    while headers and not headers[-1]:
        headers.pop()
    rows = []
    for r in values[1:]:
        cells = [cell_str(v) for v in r[: len(headers)]]
        cells += [""] * (len(headers) - len(cells))
        rows.append(cells)
    while rows and not any(rows[-1]):
        rows.pop()
    return headers, rows


def _records(values: list[list[Any]]) -> list[dict[str, str]]:
    headers, rows = _table(values)
    return [dict(zip(headers, r)) for r in rows if any(r)]


def book_from_tabs(tabs: dict[str, list[list[Any]]], *, writable: bool, source_desc: str) -> PlanBook:
    if TASKS not in tabs:
        raise ValueError(f"找不到 {TASKS} 分頁（來源：{source_desc}）")
    headers, rows = _table(tabs[TASKS])
    groups = []
    for rec in _records(tabs.get(GROUPS, [])):
        code = (rec.get("Code") or "").strip().upper()
        if code:
            groups.append({
                "code": code,
                "epic": (rec.get("Epic Name") or rec.get("Epic") or code).strip(),
                "locked": (rec.get("Locked Total (d)") or "").strip(),
                "description": (rec.get("Description") or "").strip(),
            })
    people = [
        {"name": r.get("Name", ""), "email": r.get("Email", ""),
         "account_id": r.get("Jira Account ID", ""), "role": r.get("Role", "")}
        for r in _records(tabs.get(PEOPLE, [])) if r.get("Name")
    ]
    holidays = []
    for r in _records(tabs.get(HOLIDAYS, [])):
        d = pc.parse_date_any(r.get("Date"))
        if d:
            holidays.append(d)
    settings = {}
    for r in tabs.get(SETTINGS, [])[1:]:
        if r and cell_str(r[0]):
            settings[cell_str(r[0])] = cell_str(r[1]) if len(r) > 1 else ""
    state: dict[str, Any] = {}
    for r in tabs.get(STATE, [])[1:]:
        if len(r) >= 2 and cell_str(r[0]):
            try:
                state[cell_str(r[0])] = json.loads(cell_str(r[1]))
            except json.JSONDecodeError:
                pass
    return PlanBook(headers, rows, groups, people, holidays, settings, state, writable, source_desc)


def state_to_rows(state: dict[str, Any]) -> list[list[str]]:
    rows = [["Key", "JSON"]]
    for k in sorted(state):
        rows.append([k, json.dumps(state[k], ensure_ascii=False, sort_keys=True)])
    return rows


# ---------------------------------------------------------------------------
# Template (xlsx) — 同一份範本用於本機 xlsx 與上傳成 Google Sheet
# ---------------------------------------------------------------------------

README_LINES = [
    ("專案計畫表使用說明", ""),
    ("用途", "與同事/客戶討論排程、新增工作、決定優先順序與工作內容；核准後由 ProjectSync 推到 Jira 與 Confluence。"),
    ("Tasks 分頁", "一列一個工作。ID 固定不變（PLN-<群組>-<序號>），新增列可把 ID 留空，同步時自動編號。"),
    ("草稿 / 核准", "Sync to Jira = N（或空白）是草稿/提案：只出現在 Confluence（標示「提案/討論中」），不建 Jira。改成 Y 才推到 Jira。"),
    ("優先順序", "Priority 填 P1/P2/P3，同步到 Jira 標籤 prio-P1/P2/P3。"),
    ("工期", "Duration 是工作天（Mon–Fri，扣除 Holidays 分頁的假日），0 = 里程碑。"),
    ("前置任務", "Predecessors 用 MS Project 語法：PLN-AT-01FS+1d; PLN-FX-02SS; PLN-WZ-03FF-2d。類型省略為 FS；lag 可正可負，單位 d（工作天）或 w（=5d）；多個用 ; 或 , 分隔。"),
    ("指定開始", "Start Constraint 填日期 = 不早於此日開始；填 'MSO 2026-11-02' = 必須於此日開始。"),
    ("計算欄", "Start / Finish / Jira Key / Sync Conflict / Last Sync 由程式寫入，請勿手動修改。"),
    ("討論欄", "Discussion / Comments 只在 Confluence 顯示，不送 Jira。Description / 工作內容 會送到 Jira 描述。"),
    ("刪除工作", "直接刪除整列。已建立的 Jira 單會依設定處理：預設 close（轉為完成並加標籤 plan-removed，保留 worklog），可改 delete。"),
    ("欄位主導權", "計畫欄位（名稱、描述、工期、前置、優先、群組、Sync）以 Sheet 為主；執行欄位（狀態、負責人、Jira Key）以 Jira 為主。"),
    ("雙向規則", "只有一邊修改 → 同步到另一邊；兩邊都改同一欄 → 不覆寫任何一邊，寫在 Sync Conflict 欄並在郵件報告。"),
    ("解決衝突", "在 Sync Conflict 欄輸入 use sheet（以 Sheet 為準）或 use jira（以 Jira 為準），下次同步套用後自動清空。"),
    ("Jira 拖日期", "有人在 Jira 時間軸拖曳日期 → 新開始日寫回 Start Constraint、工期依新起訖重算；仍會遵守前置相依。"),
    ("Jira 新增的工作", "在計畫 Epic 底下（或帶計畫標籤）新建的 Jira 單，會自動加成 Sheet 新列並給 ID。"),
    ("Jira 相依", "Jira 單之間新增/移除 Blocks 相依，會寫回 Predecessors（以 FS 新增）。"),
    ("Settings 分頁", "Project Start 是 T0 前的專案開始日（佔位值，請改成實際日期）。Working Days 預設 Mon–Fri。"),
]


def build_template_xlsx(
    path: Path,
    tasks: list[dict[str, str]],
    groups: list[tuple],
    people: list[tuple],
    holidays: list[tuple],
    settings: list[tuple[str, str, str]],
) -> Path:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")

    def style_header(ws, ncols: int) -> None:
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill, cell.font = header_fill, header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"

    ws_readme = wb.active
    ws_readme.title = README
    for i, (k, v) in enumerate(README_LINES, start=1):
        ws_readme.cell(row=i, column=1, value=k).font = Font(bold=True, size=14 if i == 1 else 11)
        ws_readme.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    ws_readme.column_dimensions["A"].width = 18
    ws_readme.column_dimensions["B"].width = 120

    ws = wb.create_sheet(TASKS)
    ws.append(TASK_HEADERS)
    for t in tasks:
        row = []
        for h in TASK_HEADERS:
            v = t.get(h, "")
            if h in DATE_HEADERS and v:
                v = pc.parse_date_any(v) or v
            elif h in NUMERIC_HEADERS and str(v).strip().lstrip("-").isdigit():
                v = int(v)
            row.append(v)
        ws.append(row)
    style_header(ws, len(TASK_HEADERS))
    ws.freeze_panes = "E2"
    widths = {"ID": 13, "WBS": 6, "Group": 7, "Task Name": 38, "Type": 12, "Priority": 8, "Owner": 14,
              "Duration (d)": 9, "Predecessors": 30, "Start Constraint": 12, "Start": 11, "Finish": 11,
              "Status": 12, "% Complete": 9, "Sync to Jira": 8, "Jira Key": 11,
              "Description / 工作內容": 50, "Discussion / Comments": 40, "Proposed By": 20,
              "Estimate": 8, "Notes": 36, "Sync Conflict": 36, "Last Sync": 22}
    from openpyxl.utils import get_column_letter

    letters = {h: get_column_letter(i + 1) for i, h in enumerate(TASK_HEADERS)}
    for h, w in widths.items():
        ws.column_dimensions[letters[h]].width = w
    for h, note in HEADER_NOTES.items():
        ws[f"{letters[h]}1"].comment = Comment(note, "ProjectSync")
    max_row = 1000
    for h in DATE_HEADERS:
        for r in range(2, max_row + 1):
            ws[f"{letters[h]}{r}"].number_format = "yyyy-mm-dd"
    for r in range(2, max_row + 1):
        ws[f"{letters['Predecessors']}{r}"].number_format = "@"

    def dv_list(values: list[str] | str, col: str, strict: bool = True) -> None:
        formula = values if isinstance(values, str) else '"' + ",".join(values) + '"'
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
        dv.error, dv.errorTitle = "請從下拉選單選擇", "輸入值不在清單中"
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{max_row}")

    dv_list(pc.TYPE_CHOICES, letters["Type"])
    dv_list(pc.PRIORITY_CHOICES, letters["Priority"])
    dv_list(pc.STATUS_CHOICES, letters["Status"])
    dv_list(pc.SYNC_CHOICES, letters["Sync to Jira"])
    dv_list(["Y", "N"], letters["Estimate"])
    dv_list(f"{GROUPS}!$A$2:$A$100", letters["Group"])
    dv_list(f"{PEOPLE}!$A$2:$A$200", letters["Owner"], strict=False)

    last = get_column_letter(len(TASK_HEADERS))
    rng = f"A2:{last}{max_row}"
    s, c, e = letters["Sync to Jira"], letters["Sync Conflict"], letters["Estimate"]
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f'AND($A2<>"",$D2<>"",${s}2<>"Y")'],
                         fill=PatternFill("solid", fgColor="EDEDED"), font=Font(italic=True, color="7F7F7F")))
    ws.conditional_formatting.add(
        f"{c}2:{c}{max_row}", FormulaRule(formula=[f'LEN(${c}2)>0'], fill=PatternFill("solid", fgColor="F8CBAD")))
    ws.conditional_formatting.add(
        f"{letters['Duration (d)']}2:{letters['Duration (d)']}{max_row}",
        FormulaRule(formula=[f'${e}2="Y"'], fill=PatternFill("solid", fgColor="FFF2CC")))
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f'$E2="Milestone"'], font=Font(bold=True)))

    ws_g = wb.create_sheet(GROUPS)
    ws_g.append(["Code", "Epic Name", "Locked Total (d)", "Description"])
    for g in groups:
        ws_g.append(list(g))
    style_header(ws_g, 4)
    for col, w in zip("ABCD", (8, 40, 16, 70)):
        ws_g.column_dimensions[col].width = w

    ws_p = wb.create_sheet(PEOPLE)
    ws_p.append(["Name", "Email", "Jira Account ID", "Role"])
    for p in people:
        ws_p.append(list(p))
    style_header(ws_p, 4)
    for col, w in zip("ABCD", (20, 30, 48, 20)):
        ws_p.column_dimensions[col].width = w

    ws_h = wb.create_sheet(HOLIDAYS)
    ws_h.append(["Date", "Name", "Note"])
    for d, n, note in holidays:
        ws_h.append([pc.parse_date_any(d), n, note])
    style_header(ws_h, 3)
    for r in range(2, len(holidays) + 2):
        ws_h[f"A{r}"].number_format = "yyyy-mm-dd"
    for col, w in zip("ABC", (12, 36, 20)):
        ws_h.column_dimensions[col].width = w

    ws_s = wb.create_sheet(SETTINGS)
    ws_s.append(["Key", "Value", "說明"])
    for k, v, note in settings:
        ws_s.append([k, pc.parse_date_any(v) if k == "Project Start" else v, note])
    style_header(ws_s, 3)
    ws_s["B3"].number_format = "yyyy-mm-dd"
    ws_s["B3"].fill = PatternFill("solid", fgColor="FFF2CC")
    for col, w in zip("ABC", (20, 50, 70)):
        ws_s.column_dimensions[col].width = w

    # Gantt：週欄位，條件格式畫條（Excel / Google Sheets 皆適用）
    ws_gt = wb.create_sheet(GANTT)
    ws_gt.append(["ID", "Task Name", "Start", "Finish"])
    weeks = 60
    for w in range(weeks):
        col = get_column_letter(5 + w)
        ws_gt[f"{col}1"] = f"={SETTINGS}!$B$3-WEEKDAY({SETTINGS}!$B$3,3)+{7 * w}"
        ws_gt[f"{col}1"].number_format = "mm/dd"
        ws_gt.column_dimensions[col].width = 5.5
    for r in range(2, 302):
        ws_gt[f"A{r}"] = f'=IF({TASKS}!A{r}="","",{TASKS}!A{r})'
        ws_gt[f"B{r}"] = f'=IF({TASKS}!D{r}="","",{TASKS}!D{r})'
        ws_gt[f"C{r}"] = f'=IF({TASKS}!K{r}="","",{TASKS}!K{r})'
        ws_gt[f"D{r}"] = f'=IF({TASKS}!L{r}="","",{TASKS}!L{r})'
        ws_gt[f"C{r}"].number_format = ws_gt[f"D{r}"].number_format = "yyyy-mm-dd"
    last_g = get_column_letter(4 + weeks)
    ws_gt.conditional_formatting.add(
        f"E2:{last_g}301",
        FormulaRule(formula=['AND(ISNUMBER($C2),E$1<=$D2,E$1+6>=$C2)'], fill=PatternFill("solid", fgColor="5B9BD5")))
    style_header(ws_gt, 4 + weeks)
    ws_gt.freeze_panes = "E2"
    ws_gt.column_dimensions["A"].width = 13
    ws_gt.column_dimensions["B"].width = 36
    ws_gt.column_dimensions["C"].width = ws_gt.column_dimensions["D"].width = 11

    ws_st = wb.create_sheet(STATE)
    ws_st.append(["Key", "JSON"])
    ws_st.sheet_state = "hidden"

    wb.active = wb.sheetnames.index(TASKS)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# Stores
# ---------------------------------------------------------------------------


class XlsxStore:
    """本機 xlsx（可讀寫）。Excel 開著檔案時無法寫入，會提示關閉。"""

    def __init__(self, path: Path | str):
        self.path = Path(path)

    def describe(self) -> str:
        return f"xlsx:{self.path}"

    def read(self) -> PlanBook:
        from openpyxl import load_workbook

        if not self.path.exists():
            raise FileNotFoundError(f"找不到計畫表 {self.path}")
        wb = load_workbook(self.path, data_only=True)
        tabs = {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
        return book_from_tabs(tabs, writable=True, source_desc=self.describe())

    def write(self, book: PlanBook, changes: SheetChanges) -> None:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        wb = load_workbook(self.path)
        ws = wb[TASKS]
        header_row = [canonical_header(c.value) for c in ws[1]]
        while header_row and not header_row[-1]:
            header_row.pop()

        def col_of(h: str) -> int:
            if h not in header_row:
                header_row.append(h)
                ws.cell(row=1, column=len(header_row), value=h)
            return header_row.index(h) + 1

        def put(r: int, h: str, v: Any) -> None:
            cell = ws.cell(row=r, column=col_of(h))
            if h in DATE_HEADERS and v:
                d = pc.parse_date_any(v)
                cell.value = d if d else v
                if d:
                    cell.number_format = "yyyy-mm-dd"
            elif h in NUMERIC_HEADERS and isinstance(v, str) and v.strip().lstrip("-").isdigit():
                cell.value = int(v)
            else:
                cell.value = v if v != "" else None

        for r, h, v in changes.cells:
            put(r, h, v)
        next_row = len(book.rows) + 2
        for rec in changes.appends:
            for h, v in rec.items():
                put(next_row, h, v)
            next_row += 1
        if changes.settings and SETTINGS in wb.sheetnames:
            ws_s = wb[SETTINGS]
            keys = {cell_str(ws_s.cell(row=r, column=1).value): r for r in range(2, ws_s.max_row + 1)}
            for k, v in changes.settings.items():
                r = keys.get(k) or ws_s.max_row + 1
                ws_s.cell(row=r, column=1, value=k)
                ws_s.cell(row=r, column=2, value=v)
        if changes.state is not None:
            ws_st = wb[STATE] if STATE in wb.sheetnames else wb.create_sheet(STATE)
            ws_st.sheet_state = "hidden"
            ws_st.delete_rows(1, ws_st.max_row)
            for row in state_to_rows(changes.state):
                ws_st.append(row)
        _ = get_column_letter
        try:
            wb.save(self.path)
        except PermissionError as exc:
            raise RuntimeError(f"無法寫入 {self.path}（檔案可能正在 Excel 中開啟，請先關閉）") from exc


def google_sheet_id(url_or_id: str) -> str:
    m = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", url_or_id or "")
    return m.group(1) if m else (url_or_id or "").strip()


class ExportStore:
    """Google 公開匯出（唯讀）。"""

    def __init__(self, url_or_id: str):
        self.sheet_id = google_sheet_id(url_or_id)

    def describe(self) -> str:
        return f"google-export:{self.sheet_id}"

    def read(self) -> PlanBook:
        from openpyxl import load_workbook

        url = f"https://docs.google.com/spreadsheets/d/{self.sheet_id}/export?format=xlsx"
        r = requests.get(url, timeout=60)
        if r.status_code != 200 or not r.content.startswith(b"PK"):
            raise RuntimeError(
                f"下載 Google Sheet 失敗（HTTP {r.status_code}）。請確認共用設定為「知道連結的任何人可檢視」，"
                "或設定 service account 改用 google 模式。")
        wb = load_workbook(io.BytesIO(r.content), data_only=True)
        tabs = {ws.title: [list(x) for x in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
        return book_from_tabs(tabs, writable=False, source_desc=self.describe())

    def write(self, book: PlanBook, changes: SheetChanges) -> None:
        raise RuntimeError("Google 公開匯出為唯讀，無法回寫")


SCOPES = ["https://www.googleapis.com/auth/spreadsheets", "https://www.googleapis.com/auth/drive"]


def google_credentials(sa_json_path: str):
    from google.oauth2.service_account import Credentials

    return Credentials.from_service_account_file(sa_json_path, scopes=SCOPES)


class GoogleStore:
    """Google Sheets（service account，可讀寫）。"""

    def __init__(self, url_or_id: str, sa_json_path: str):
        import gspread

        self.sheet_id = google_sheet_id(url_or_id)
        self.gc = gspread.authorize(google_credentials(sa_json_path))
        self.sh = self.gc.open_by_key(self.sheet_id)

    def describe(self) -> str:
        return f"google:{self.sheet_id}"

    def read(self) -> PlanBook:
        tabs = {}
        for ws in self.sh.worksheets():
            tabs[ws.title] = ws.get_all_values(value_render_option="FORMATTED_VALUE")
        return book_from_tabs(tabs, writable=True, source_desc=self.describe())

    def write(self, book: PlanBook, changes: SheetChanges) -> None:
        import gspread
        from gspread.utils import rowcol_to_a1

        ws = self.sh.worksheet(TASKS)
        headers = list(book.headers)
        new_headers = []

        def col_of(h: str) -> int:
            if h not in headers:
                headers.append(h)
                new_headers.append(h)
            return headers.index(h) + 1

        raw, entered = [], []
        for r, h, v in changes.cells:
            target = entered if h in DATE_HEADERS or h in NUMERIC_HEADERS else raw
            target.append({"range": rowcol_to_a1(r, col_of(h)), "values": [[v if v is not None else ""]]})
        next_row = len(book.rows) + 2
        for rec in changes.appends:
            for h, v in rec.items():
                target = entered if h in DATE_HEADERS or h in NUMERIC_HEADERS else raw
                target.append({"range": rowcol_to_a1(next_row, col_of(h)), "values": [[v]]})
            next_row += 1
        need_rows = max(ws.row_count, next_row + 5)
        need_cols = max(ws.col_count, len(headers))
        if need_rows > ws.row_count or need_cols > ws.col_count:
            ws.resize(rows=need_rows, cols=need_cols)
        for h in new_headers:
            raw.append({"range": rowcol_to_a1(1, headers.index(h) + 1), "values": [[h]]})
        if raw:
            ws.batch_update(raw, value_input_option="RAW")
        if entered:
            ws.batch_update(entered, value_input_option="USER_ENTERED")
        if changes.settings:
            try:
                ws_s = self.sh.worksheet(SETTINGS)
                vals = ws_s.get_all_values()
                keys = {row[0]: i + 1 for i, row in enumerate(vals) if row}
                upd = []
                nxt = len(vals) + 1
                for k, v in changes.settings.items():
                    r = keys.get(k)
                    if not r:
                        r, nxt = nxt, nxt + 1
                    upd.append({"range": f"A{r}:B{r}", "values": [[k, v]]})
                ws_s.batch_update(upd, value_input_option="USER_ENTERED")
            except gspread.WorksheetNotFound:
                pass
        if changes.state is not None:
            try:
                ws_st = self.sh.worksheet(STATE)
            except gspread.WorksheetNotFound:
                ws_st = self.sh.add_worksheet(STATE, rows=500, cols=2)
            rows = state_to_rows(changes.state)
            ws_st.clear()
            if ws_st.row_count < len(rows) + 5:
                ws_st.resize(rows=len(rows) + 50, cols=2)
            ws_st.update(values=rows, range_name="A1", value_input_option="RAW")


def upload_xlsx_as_google_sheet(xlsx: Path, title: str, sa_json_path: str, share_with: list[str]) -> str:
    """把範本 xlsx 上傳並轉成 Google Sheet，分享給指定帳號（Editor）。回傳 sheet URL。

    service account 自己沒有 Drive 空間時（新式 SA 常見）會失敗：此時請使用者自建空白 Sheet、
    分享給 SA 後改用「匯入到既有 Sheet」流程（GUI 精靈會提示）。
    """
    from google.auth.transport.requests import AuthorizedSession

    sess = AuthorizedSession(google_credentials(sa_json_path))
    meta = {"name": title, "mimeType": "application/vnd.google-apps.spreadsheet"}
    boundary = "projectsync_boundary"
    body = (
        f"--{boundary}\r\nContent-Type: application/json; charset=UTF-8\r\n\r\n{json.dumps(meta)}\r\n"
        f"--{boundary}\r\nContent-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n"
    ).encode() + xlsx.read_bytes() + f"\r\n--{boundary}--".encode()
    r = sess.post(
        "https://www.googleapis.com/upload/drive/v3/files?uploadType=multipart&supportsAllDrives=true",
        data=body, headers={"Content-Type": f"multipart/related; boundary={boundary}"}, timeout=120)
    if not r.ok:
        raise RuntimeError(f"上傳 Google Sheet 失敗：{r.status_code} {r.text[:300]}")
    file_id = r.json()["id"]
    for email in share_with:
        p = sess.post(
            f"https://www.googleapis.com/drive/v3/files/{file_id}/permissions?sendNotificationEmail=false",
            json={"type": "user", "role": "writer", "emailAddress": email}, timeout=60)
        if not p.ok:
            raise RuntimeError(f"分享給 {email} 失敗：{p.status_code} {p.text[:200]}")
    return f"https://docs.google.com/spreadsheets/d/{file_id}/edit"


def import_xlsx_into_existing_sheet(xlsx: Path, sheet_url: str, sa_json_path: str) -> str:
    """把範本內容覆蓋寫進使用者自建、已分享給 SA 的空白 Sheet（Drive files.update 轉換）。"""
    from google.auth.transport.requests import AuthorizedSession

    sess = AuthorizedSession(google_credentials(sa_json_path))
    file_id = google_sheet_id(sheet_url)
    r = sess.patch(
        f"https://www.googleapis.com/upload/drive/v3/files/{file_id}?uploadType=media&supportsAllDrives=true",
        data=xlsx.read_bytes(),
        headers={"Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
        timeout=120)
    if not r.ok:
        raise RuntimeError(f"匯入既有 Sheet 失敗：{r.status_code} {r.text[:300]}")
    return f"https://docs.google.com/spreadsheets/d/{file_id}/edit"
