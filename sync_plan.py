#!/usr/bin/env python3
"""
C27 Project Plan：Google Sheet（計畫主檔） <-> Jira PMWC（雙向） -> Confluence 頁面。

流程（--direction both，預設）：
  1. 讀 Google Sheet（Tasks / Groups / People / Holidays / Settings / _sync_state）
  2. 讀 Jira 上所有 label=c27-plan 的計畫單（與計畫 Epic 底下的子單）
  3. 依「上次同步快照」三方比對每個欄位：只有一邊改 → 同步到另一邊；兩邊都改 → 衝突不覆寫
  4. 套用 Jira 端的變更到 Sheet（日期拖曳 → Start Constraint/工期；新相依 → Predecessors；
     Jira 新建的計畫單 → 新增列）
  5. 重新排程（工作天 CPM forward pass，FS/SS/FF/SF + lag、限制日、假日）
  6. 推送到 Jira（Epic / 任務 / Start / Due / Blocks 相依）並寫回 Sheet（Start/Finish/Jira Key…）
  7. 更新 Confluence「C27 Project Plan」頁面，寄送結果郵件

只會動 label = c27-plan 的 Jira 單；Register（sync_register.py）的單一律不碰。
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import html
import io
import json
import os
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Callable

import requests
import yaml

import sync_register as sr

SCRIPT_DIR = Path(__file__).resolve().parent
TW_TZ = timezone(timedelta(hours=8))

PLAN_ID_RE = re.compile(r"^PLN-[A-Z][A-Z0-9]*-\d{2,3}$")
PLAN_SUMMARY_RE = re.compile(r"^(PLN-[A-Z][A-Z0-9]*-\d{2,3})_(.*)$", re.DOTALL)
REGISTER_PAGE_IDS = {"189300768"}
PROTECTED_PAGE_TITLES = {"mail_checking"}
DESC_MARKER = "［C27 Plan Sync］"

TASKS_TAB = "Tasks"
GROUPS_TAB = "Groups"
PEOPLE_TAB = "People"
HOLIDAYS_TAB = "Holidays"
SETTINGS_TAB = "Settings"
README_TAB = "說明"
GANTT_TAB = "Gantt"
STATE_TAB = "_sync_state"

# canonical key -> accepted header names (case-insensitive)
TASK_COLUMNS: dict[str, list[str]] = {
    "id": ["ID"],
    "wbs": ["WBS"],
    "group": ["Group", "Group/Epic", "Epic"],
    "name": ["Task Name", "Task", "Name"],
    "type": ["Type"],
    "priority": ["Priority"],
    "owner": ["Owner"],
    "duration": ["Duration (d)", "Duration"],
    "preds": ["Predecessors"],
    "constraint": ["Start Constraint"],
    "start": ["Start"],
    "finish": ["Finish"],
    "status": ["Status"],
    "pct": ["% Complete"],
    "sync": ["Sync to Jira", "Sync"],
    "jira_key": ["Jira Key"],
    "description": ["Description / 工作內容", "Description"],
    "discussion": ["Discussion / Comments", "Discussion"],
    "proposed_by": ["Proposed By"],
    "estimate": ["Estimate"],
    "notes": ["Notes"],
    "conflict": ["Sync Conflict"],
    "last_sync": ["Last Sync"],
}
DATE_COLUMNS = {"start", "finish", "constraint"}
NUMERIC_COLUMNS = {"duration", "pct"}

TYPE_TASK, TYPE_MILESTONE, TYPE_PROCUREMENT = "Task", "Milestone", "Procurement"
STATUS_OPTIONS = ["Not Started", "In Progress", "Blocked", "Done", "Cancelled"]
PRIORITY_OPTIONS = ["P1", "P2", "P3"]

JIRA_TODO, JIRA_DOING, JIRA_DONE = "待辦事項", "進行中", "完成"
SHEET_TO_JIRA_STATUS = {
    "Not Started": JIRA_TODO,
    "In Progress": JIRA_DOING,
    "Blocked": JIRA_DOING,
    "Done": JIRA_DONE,
    "Cancelled": JIRA_DONE,
}
JIRA_STATUS_ALIASES = {
    "to do": JIRA_TODO, "todo": JIRA_TODO, "待辦": JIRA_TODO, "待辦事項": JIRA_TODO,
    "in progress": JIRA_DOING, "進行中": JIRA_DOING,
    "done": JIRA_DONE, "完成": JIRA_DONE, "closed": JIRA_DONE,
}

SYNC_FIELDS = ["name", "description", "priority", "group", "status", "owner", "start", "finish", "preds"]
FIELD_MASTER = {
    "name": "sheet", "description": "sheet", "priority": "sheet", "group": "sheet",
    "start": "sheet", "finish": "sheet", "preds": "sheet",
    "status": "jira", "owner": "jira",
}
FIELD_LABELS = {
    "name": "Task Name", "description": "Description", "priority": "Priority",
    "group": "Group", "status": "Status", "owner": "Owner", "start": "Start",
    "finish": "Finish", "preds": "Predecessors",
}


class PlanError(Exception):
    """設定或資料錯誤（會中止同步，不做任何寫入）。"""


class ScheduleError(PlanError):
    pass


# ---------------------------------------------------------------------------
# Working-day calendar
# ---------------------------------------------------------------------------

WEEKDAY_NAMES = {
    "mon": 0, "tue": 1, "wed": 2, "thu": 3, "fri": 4, "sat": 5, "sun": 6,
    "一": 0, "二": 1, "三": 2, "四": 3, "五": 4, "六": 5, "日": 6,
}


def parse_working_days(text: str) -> set[int]:
    raw = (text or "").strip()
    if not raw:
        return {0, 1, 2, 3, 4}
    days: set[int] = set()
    for tok in re.split(r"[,\s;、/]+", raw):
        tok = tok.strip().lower()
        if not tok:
            continue
        m = re.match(r"^([a-z]{3}|[一二三四五六日])[a-z]*\s*[-~–]\s*([a-z]{3}|[一二三四五六日])", tok)
        if m and m.group(1)[:3] in WEEKDAY_NAMES and m.group(2)[:3] in WEEKDAY_NAMES:
            a, b = WEEKDAY_NAMES[m.group(1)[:3]], WEEKDAY_NAMES[m.group(2)[:3]]
            days.update(range(a, b + 1))
            continue
        key = tok[:3] if tok[:1].isascii() else tok[:1]
        if key in WEEKDAY_NAMES:
            days.add(WEEKDAY_NAMES[key])
    if not days:
        raise PlanError(f"Working Days 設定無法解析：{text!r}")
    return days


class WorkCalendar:
    """工作天索引：index 0 = 專案開始日（若非工作日則順延）。"""

    def __init__(self, start: date, workdays: set[int] | None = None, holidays: set[date] | None = None):
        self.workdays = set(workdays if workdays is not None else {0, 1, 2, 3, 4})
        if not self.workdays:
            raise PlanError("工作日設定為空")
        self.holidays = set(holidays or ())
        self.base = self.roll_forward(start)
        self._fwd: list[date] = [self.base]
        self._back: list[date] = []

    def is_workday(self, d: date) -> bool:
        return d.weekday() in self.workdays and d not in self.holidays

    def roll_forward(self, d: date) -> date:
        for _ in range(3700):
            if self.is_workday(d):
                return d
            d += timedelta(days=1)
        raise PlanError("找不到工作日（假日設定異常？）")

    def _step(self, d: date, delta: int) -> date:
        for _ in range(3700):
            d += timedelta(days=delta)
            if self.is_workday(d):
                return d
        raise PlanError("找不到工作日（假日設定異常？）")

    def date_of(self, idx: int) -> date:
        if idx >= 0:
            while len(self._fwd) <= idx:
                self._fwd.append(self._step(self._fwd[-1], 1))
            return self._fwd[idx]
        need = -idx
        while len(self._back) < need:
            prev = self._back[-1] if self._back else self.base
            self._back.append(self._step(prev, -1))
        return self._back[need - 1]

    def index_of(self, d: date) -> int:
        d = self.roll_forward(d)
        if d >= self.base:
            while self._fwd[-1] < d:
                self._fwd.append(self._step(self._fwd[-1], 1))
            lo, hi = 0, len(self._fwd) - 1
            while lo < hi:
                mid = (lo + hi) // 2
                if self._fwd[mid] < d:
                    lo = mid + 1
                else:
                    hi = mid
            return lo
        i = -1
        while self.date_of(i) > d:
            i -= 1
        return i

    def workdays_between(self, start: date, finish: date) -> int:
        """含頭含尾的工作天數；finish < start 時回傳 0。"""
        if finish < start:
            return 0
        return self.index_of(finish + timedelta(days=1)) - self.index_of(start)


# ---------------------------------------------------------------------------
# Predecessor parsing (Microsoft Project style)
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class Pred:
    id: str
    kind: str = "FS"
    lag: int = 0

    def format(self) -> str:
        if self.kind == "FS" and self.lag == 0:
            return self.id
        lag = f"{self.lag:+d}d" if self.lag else ""
        return f"{self.id}{self.kind}{lag}"


_PRED_SPLIT_RE = re.compile(r"[;,，；\n]+")
_PRED_REST_RE = re.compile(
    r"^\s*(FS|SS|FF|SF)?\s*(?:([+-])\s*(\d+(?:\.\d+)?)\s*"
    r"(d|day|days|wd|w|wk|wks|week|weeks|工作天|天|週|周)?)?\s*$",
    re.IGNORECASE,
)
_PRED_FALLBACK_RE = re.compile(
    r"^(?P<id>.+?)(?P<rest>(?:FS|SS|FF|SF)?\s*(?:[+-]\s*\d+(?:\.\d+)?\s*"
    r"(?:d|day|days|wd|w|wk|wks|week|weeks|工作天|天|週|周)?)?)\s*$",
    re.IGNORECASE,
)


def _parse_pred_rest(rest: str) -> tuple[str, int] | None:
    m = _PRED_REST_RE.match(rest or "")
    if not m:
        return None
    kind = (m.group(1) or "FS").upper()
    lag = 0
    if m.group(3):
        value = float(m.group(3))
        unit = (m.group(4) or "d").lower()
        if unit in ("w", "wk", "wks", "week", "weeks", "週", "周"):
            value *= 5
        if abs(value - round(value)) > 1e-9:
            return None
        lag = int(round(value))
        if m.group(2) == "-":
            lag = -lag
    return kind, lag


def parse_predecessors(text: str, known_ids: set[str] | None = None) -> tuple[list[Pred], list[str]]:
    """解析 `PLN-AT-01FS+1d; PLN-FX-02SS, PLN-WZ-03FF-2d`。回傳 (preds, errors)。"""
    preds: list[Pred] = []
    errors: list[str] = []
    known_sorted = sorted(known_ids or (), key=len, reverse=True)
    for raw in _PRED_SPLIT_RE.split(text or ""):
        tok = raw.strip()
        if not tok:
            continue
        pid, rest = "", ""
        if known_sorted:
            up = tok.upper()
            for k in known_sorted:
                if up.startswith(k.upper()):
                    pid, rest = k, tok[len(k):]
                    break
            if not pid:
                m = _PRED_FALLBACK_RE.match(tok)
                guess = m.group("id").strip() if m else tok
                errors.append(f"未知前置任務 {guess!r}")
                continue
        else:
            m = _PRED_FALLBACK_RE.match(tok)
            if not m:
                errors.append(f"前置任務語法錯誤 {tok!r}")
                continue
            pid, rest = m.group("id").strip().upper(), m.group("rest")
        parsed = _parse_pred_rest(rest)
        if parsed is None:
            errors.append(f"前置任務語法錯誤 {tok!r}（例：PLN-AT-01FS+1d）")
            continue
        kind, lag = parsed
        pred = Pred(pid, kind, lag)
        if pred not in preds:
            preds.append(pred)
    return preds, errors


def format_predecessors(preds: list[Pred]) -> str:
    return "; ".join(p.format() for p in preds)


# ---------------------------------------------------------------------------
# Value normalization helpers
# ---------------------------------------------------------------------------

def norm_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def norm_multiline(value: Any) -> str:
    lines = [ln.rstrip() for ln in norm_text(value).split("\n")]
    while lines and not lines[-1]:
        lines.pop()
    while lines and not lines[0]:
        lines.pop(0)
    return "\n".join(lines)


def parse_date_any(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        if 20000 < float(value) < 80000:
            return date(1899, 12, 30) + timedelta(days=int(value))
        return None
    text = str(value).strip()
    if not text:
        return None
    m = re.search(r"(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})", text)
    if m:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.search(r"\b(\d{1,2})/(\d{1,2})/(\d{4})\b", text)
    if m:
        return date(int(m.group(3)), int(m.group(1)), int(m.group(2)))
    if re.fullmatch(r"\d{5}(\.\d+)?", text):
        return parse_date_any(float(text))
    return None


def normalize_type(value: str, duration: int | None) -> str:
    v = (value or "").strip().lower()
    if v in ("milestone", "里程碑", "ms") or duration == 0:
        return TYPE_MILESTONE
    if v in ("procurement", "採購", "purchase", "pr"):
        return TYPE_PROCUREMENT
    return TYPE_TASK


def normalize_priority(value: str) -> str:
    v = (value or "").strip().upper()
    mapping = {
        "P1": "P1", "HIGH": "P1", "H": "P1", "高": "P1", "1": "P1", "P0": "P1",
        "P2": "P2", "MEDIUM": "P2", "MED": "P2", "M": "P2", "中": "P2", "2": "P2",
        "P3": "P3", "LOW": "P3", "L": "P3", "低": "P3", "3": "P3",
    }
    return mapping.get(v, "")


def normalize_status(value: str) -> str:
    v = (value or "").strip().lower()
    mapping = {
        "": "Not Started", "not started": "Not Started", "未開始": "Not Started", "todo": "Not Started",
        "to do": "Not Started", "待辦": "Not Started", "待辦事項": "Not Started", "open": "Not Started",
        "in progress": "In Progress", "進行中": "In Progress", "wip": "In Progress", "ongoing": "In Progress",
        "blocked": "Blocked", "卡關": "Blocked", "暫停": "Blocked", "on hold": "Blocked",
        "done": "Done", "完成": "Done", "closed": "Done", "completed": "Done", "已完成": "Done",
        "cancelled": "Cancelled", "canceled": "Cancelled", "取消": "Cancelled",
    }
    return mapping.get(v, "Not Started")


def parse_sync_flag(value: str) -> bool:
    return (value or "").strip().lower() in ("y", "yes", "true", "1", "v", "✓", "✔", "是", "x")


def normalize_jira_status(name: str) -> str:
    v = (name or "").strip()
    return JIRA_STATUS_ALIASES.get(v.lower(), v)


def parse_constraint(text: str) -> tuple[str, date | None]:
    """'2026-11-02' / 'SNET 2026-11-02' / 'MSO 2026-11-02' / '不早於 2026-11-02'."""
    raw = norm_text(text)
    if not raw:
        return "", None
    d = parse_date_any(raw)
    if d is None:
        return "?", None
    upper = raw.upper()
    if upper.startswith("MSO") or "指定" in raw or "必須" in raw:
        return "MSO", d
    return "SNET", d


# ---------------------------------------------------------------------------
# Plan model
# ---------------------------------------------------------------------------

@dataclass
class PlanTask:
    row: int  # sheet row number (1-based); 0 = new row to append
    id: str
    cells: dict[str, str]
    wbs: str = ""
    group: str = ""
    name: str = ""
    type: str = TYPE_TASK
    priority: str = ""
    owner: str = ""
    duration: int = 1
    preds_text: str = ""
    constraint_text: str = ""
    status: str = "Not Started"
    sync: bool = False
    jira_key: str = ""
    description: str = ""
    discussion: str = ""
    estimate: bool = False
    conflict_text: str = ""
    preds: list[Pred] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    start_idx: int = 0
    end_idx: int = 0
    start: date | None = None
    finish: date | None = None

    @property
    def is_milestone(self) -> bool:
        return self.duration == 0

    def resolution(self) -> str:
        t = (self.conflict_text or "").strip().lower()
        if t.startswith("use sheet") or t.startswith("以sheet") or t.startswith("以 sheet"):
            return "sheet"
        if t.startswith("use jira") or t.startswith("以jira") or t.startswith("以 jira"):
            return "jira"
        return ""


@dataclass
class Group:
    code: str
    epic_name: str
    locked_total: int | None = None
    description: str = ""


@dataclass
class Person:
    name: str
    email: str = ""
    account_id: str = ""


@dataclass
class PlanData:
    task_headers: list[str]
    tasks: list[PlanTask]
    groups: dict[str, Group]
    people: list[Person]
    holidays: set[date]
    settings: dict[str, str]
    state: dict[str, Any]
    task_row_count: int = 1  # rows used incl. header

    def task_by_id(self) -> dict[str, PlanTask]:
        return {t.id: t for t in self.tasks if t.id}


def header_index(headers: list[str]) -> dict[str, int]:
    norm = {re.sub(r"\s+", " ", h.strip().lower()): i for i, h in enumerate(headers) if h and h.strip()}
    out: dict[str, int] = {}
    for key, names in TASK_COLUMNS.items():
        for n in names:
            i = norm.get(n.lower())
            if i is not None:
                out[key] = i
                break
    return out


def parse_int(value: Any, default: int) -> int:
    text = norm_text(value)
    if not text:
        return default
    try:
        return int(round(float(text.replace("d", "").replace("天", "").strip())))
    except ValueError:
        return default


def parse_tasks(values: list[list[Any]]) -> tuple[list[str], list[PlanTask]]:
    if not values:
        raise PlanError("Tasks 分頁是空的")
    headers = [norm_text(h) for h in values[0]]
    idx = header_index(headers)
    for required in ("id", "name", "duration", "preds"):
        if required not in idx:
            raise PlanError(f"Tasks 分頁缺少欄位 {TASK_COLUMNS[required][0]!r}")
    tasks: list[PlanTask] = []
    for r, raw in enumerate(values[1:], start=2):
        cells = {k: norm_text(raw[i]) if i < len(raw) else "" for k, i in idx.items()}
        for k in DATE_COLUMNS:
            if k in idx and idx[k] < len(raw):
                d = parse_date_any(raw[idx[k]])
                if d:
                    cells[k] = d.isoformat()
        if not cells.get("id") and not cells.get("name"):
            continue
        dur = max(0, parse_int(cells.get("duration"), 1))
        t = PlanTask(
            row=r,
            id=cells.get("id", "").strip().upper(),
            cells=cells,
            wbs=cells.get("wbs", ""),
            group=cells.get("group", "").strip().upper(),
            name=cells.get("name", ""),
            type=normalize_type(cells.get("type", ""), dur),
            priority=normalize_priority(cells.get("priority", "")),
            owner=cells.get("owner", ""),
            duration=dur,
            preds_text=cells.get("preds", ""),
            constraint_text=cells.get("constraint", ""),
            status=normalize_status(cells.get("status", "")),
            sync=parse_sync_flag(cells.get("sync", "")),
            jira_key=cells.get("jira_key", "").strip().upper(),
            description=norm_multiline(cells.get("description", "")),
            discussion=norm_multiline(cells.get("discussion", "")),
            estimate=parse_sync_flag(cells.get("estimate", "")),
            conflict_text=cells.get("conflict", ""),
        )
        if t.type == TYPE_MILESTONE:
            t.duration = 0
        if t.jira_key and not sr.is_valid_jira_key(t.jira_key):
            t.errors.append(f"Jira Key 格式錯誤 {t.jira_key!r}")
            t.jira_key = ""
        tasks.append(t)
    return headers, tasks


def next_plan_id(group: str, existing: set[str]) -> str:
    code = re.sub(r"[^A-Z0-9]", "", (group or "").upper()) or "NEW"
    if not code[0].isalpha():
        code = "N" + code
    n = 1
    while True:
        cand = f"PLN-{code}-{n:02d}"
        if cand not in existing:
            return cand
        n += 1


def assign_missing_ids(tasks: list[PlanTask]) -> list[PlanTask]:
    existing = {t.id for t in tasks if t.id}
    assigned = []
    for t in tasks:
        if not t.id:
            t.id = next_plan_id(t.group, existing)
            existing.add(t.id)
            assigned.append(t)
    return assigned


def validate_ids(tasks: list[PlanTask]) -> None:
    seen: dict[str, int] = {}
    problems = []
    for t in tasks:
        if not PLAN_ID_RE.match(t.id):
            problems.append(f"列 {t.row}: ID {t.id!r} 不符合格式 PLN-<群組>-<數字>（例 PLN-AT-01）")
        if sr.REGISTER_ID_RE.match(t.id):
            problems.append(f"列 {t.row}: ID {t.id!r} 會與 Register ID 格式衝突")
        if t.id in seen:
            problems.append(f"ID 重複：{t.id}（列 {seen[t.id]} 與列 {t.row}）")
        seen[t.id] = t.row
    if problems:
        raise PlanError("Tasks 分頁 ID 錯誤：\n  " + "\n  ".join(problems))


def resolve_predecessors(tasks: list[PlanTask]) -> None:
    known = {t.id for t in tasks}
    for t in tasks:
        preds, errors = parse_predecessors(t.preds_text, known)
        clean = []
        for p in preds:
            if p.id == t.id:
                errors.append("前置任務不可指向自己")
                continue
            clean.append(p)
        t.preds = clean
        t.errors.extend(errors)


# ---------------------------------------------------------------------------
# Scheduler (CPM forward pass on working-day indices; task occupies [start, end))
# ---------------------------------------------------------------------------

def topo_order(tasks: list[PlanTask]) -> list[PlanTask]:
    by_id = {t.id: t for t in tasks}
    indeg = {t.id: 0 for t in tasks}
    succs: dict[str, list[str]] = {t.id: [] for t in tasks}
    for t in tasks:
        for p in t.preds:
            if p.id in by_id:
                indeg[t.id] += 1
                succs[p.id].append(t.id)
    ready = [t.id for t in tasks if indeg[t.id] == 0]
    order: list[str] = []
    while ready:
        cur = ready.pop(0)
        order.append(cur)
        for s in succs[cur]:
            indeg[s] -= 1
            if indeg[s] == 0:
                ready.append(s)
    if len(order) != len(tasks):
        remaining = {tid for tid in indeg if tid not in set(order)}
        raise ScheduleError("偵測到循環相依：" + find_cycle(remaining, by_id))
    return [by_id[i] for i in order]


def find_cycle(remaining: set[str], by_id: dict[str, PlanTask]) -> str:
    color: dict[str, int] = {}
    stack: list[str] = []

    def dfs(node: str) -> list[str] | None:
        color[node] = 1
        stack.append(node)
        for p in by_id[node].preds:
            if p.id not in remaining:
                continue
            if color.get(p.id) == 1:
                return stack[stack.index(p.id):] + [p.id]
            if color.get(p.id) is None:
                found = dfs(p.id)
                if found:
                    return found
        stack.pop()
        color[node] = 2
        return None

    for n in sorted(remaining):
        if color.get(n) is None:
            cyc = dfs(n)
            if cyc:
                return " ← ".join(cyc) + "（箭頭方向：任務 ← 其前置任務）"
    return ", ".join(sorted(remaining))


def compute_schedule(tasks: list[PlanTask], cal: WorkCalendar) -> list[str]:
    """計算每個任務的 Start/Finish；回傳警告。循環相依時丟 ScheduleError。"""
    warnings: list[str] = []
    by_id = {t.id: t for t in tasks}
    for t in topo_order(tasks):
        dur = t.duration
        pred_bound: int | None = None
        shift_display = False
        for p in t.preds:
            P = by_id[p.id]
            if p.kind == "FS":
                b, from_end = P.end_idx + p.lag, True
            elif p.kind == "SS":
                b, from_end = P.start_idx + p.lag, False
            elif p.kind == "FF":
                b, from_end = P.end_idx + p.lag - dur, True
            else:  # SF
                b, from_end = P.start_idx + p.lag - dur, False
            if pred_bound is None or b > pred_bound:
                pred_bound, shift_display = b, from_end
            elif b == pred_bound and from_end:
                shift_display = True
        kind, cdate = parse_constraint(t.constraint_text)
        if kind == "?":
            t.errors.append(f"Start Constraint 無法解析 {t.constraint_text!r}（例：2026-11-02 或 MSO 2026-11-02）")
            kind = ""
        cidx = cal.index_of(cdate) if cdate else None
        if kind == "MSO" and cidx is not None:
            start = cidx
            if pred_bound is not None and pred_bound > cidx:
                warnings.append(f"{t.id}: MSO {cdate} 早於前置任務允許的開始日（相依關係被違反）")
            driven = False
        else:
            lower = max(0, cidx) if cidx is not None else 0
            start = lower if pred_bound is None else max(lower, pred_bound)
            driven = pred_bound is not None and pred_bound >= lower and pred_bound == start
        if start < 0:
            start = 0
        t.start_idx, t.end_idx = start, start + dur
        if dur > 0:
            t.start = cal.date_of(start)
            t.finish = cal.date_of(start + dur - 1)
        else:
            disp = start - 1 if (driven and shift_display and start > 0) else start
            t.start = t.finish = cal.date_of(disp)
    return warnings


def validate_group_totals(tasks: list[PlanTask], groups: dict[str, Group]) -> list[tuple[str, int, int, bool]]:
    results = []
    for code, g in groups.items():
        if g.locked_total is None:
            continue
        actual = sum(
            t.duration for t in tasks
            if t.group == code and t.type == TYPE_TASK and t.status != "Cancelled"
        )
        results.append((code, g.locked_total, actual, actual == g.locked_total))
    return results


def group_spans(tasks: list[PlanTask]) -> dict[str, tuple[date, date]]:
    spans: dict[str, tuple[date, date]] = {}
    for t in tasks:
        if not t.start or not t.finish:
            continue
        s, f = spans.get(t.group, (t.start, t.finish))
        spans[t.group] = (min(s, t.start), max(f, t.finish))
    return spans


# ---------------------------------------------------------------------------
# Sheet parsing of auxiliary tabs
# ---------------------------------------------------------------------------

def _rows_as_dicts(values: list[list[Any]]) -> list[dict[str, str]]:
    if not values:
        return []
    headers = [norm_text(h).lower() for h in values[0]]
    out = []
    for raw in values[1:]:
        row = {h: norm_text(raw[i]) if i < len(raw) else "" for i, h in enumerate(headers) if h}
        if any(row.values()):
            out.append(row)
    return out


def _pick(row: dict[str, str], *names: str) -> str:
    for n in names:
        for k, v in row.items():
            if k.startswith(n.lower()):
                return v
    return ""


def parse_groups(values: list[list[Any]]) -> dict[str, Group]:
    groups: dict[str, Group] = {}
    for row in _rows_as_dicts(values):
        code = _pick(row, "code").strip().upper()
        if not code:
            continue
        locked = _pick(row, "locked")
        groups[code] = Group(
            code=code,
            epic_name=_pick(row, "epic") or f"PLAN-{code}",
            locked_total=parse_int(locked, 0) if locked else None,
            description=_pick(row, "description", "說明"),
        )
    return groups


def parse_people(values: list[list[Any]]) -> list[Person]:
    people = []
    for row in _rows_as_dicts(values):
        name = _pick(row, "name")
        if name:
            people.append(Person(name, _pick(row, "email"), _pick(row, "jira account", "account")))
    return people


def parse_holidays(values: list[list[Any]]) -> set[date]:
    out: set[date] = set()
    for raw in values[1:] if values else []:
        if raw:
            d = parse_date_any(raw[0])
            if d:
                out.add(d)
    return out


def parse_settings(values: list[list[Any]]) -> dict[str, str]:
    out: dict[str, str] = {}
    for raw in values[1:] if values else []:
        if len(raw) >= 2 and norm_text(raw[0]):
            key = norm_text(raw[0]).lower()
            val = raw[1]
            d = parse_date_any(val) if key.startswith("project start") else None
            out[key] = d.isoformat() if d else norm_text(val)
    return out


def parse_state(values: list[list[Any]]) -> dict[str, Any]:
    state: dict[str, Any] = {"meta": {}, "items": {}}
    for raw in values[1:] if values else []:
        if len(raw) < 2 or not norm_text(raw[0]):
            continue
        try:
            data = json.loads(str(raw[1]))
        except (TypeError, ValueError):
            continue
        key = norm_text(raw[0])
        if key == "__meta__":
            state["meta"] = data
        else:
            state["items"][key] = data
    return state


def build_plan_data(tabs: dict[str, list[list[Any]]]) -> PlanData:
    headers, tasks = parse_tasks(tabs.get(TASKS_TAB) or [])
    data = PlanData(
        task_headers=headers,
        tasks=tasks,
        groups=parse_groups(tabs.get(GROUPS_TAB) or []),
        people=parse_people(tabs.get(PEOPLE_TAB) or []),
        holidays=parse_holidays(tabs.get(HOLIDAYS_TAB) or []),
        settings=parse_settings(tabs.get(SETTINGS_TAB) or []),
        state=parse_state(tabs.get(STATE_TAB) or []),
        task_row_count=len(tabs.get(TASKS_TAB) or []),
    )
    return data


def project_calendar(data: PlanData) -> WorkCalendar:
    start = parse_date_any(data.settings.get("project start", ""))
    if not start:
        raise PlanError("Settings 分頁缺少 Project Start（專案開始日）")
    return WorkCalendar(start, parse_working_days(data.settings.get("working days", "")), data.holidays)


# ---------------------------------------------------------------------------
# Sheet backends
# ---------------------------------------------------------------------------

def google_credentials_json() -> str:
    return (os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON") or os.environ.get("GOOGLE_SHEETS_JSON") or "").strip()


def gspread_client():
    import gspread

    raw = google_credentials_json()
    if not raw:
        return None
    info = json.loads(raw)
    return gspread.service_account_from_dict(info)


def service_account_email() -> str:
    raw = google_credentials_json()
    if not raw:
        return ""
    try:
        return str(json.loads(raw).get("client_email") or "")
    except ValueError:
        return ""


def sheet_id_from_url(url: str) -> str:
    m = re.search(r"/spreadsheets/d/([a-zA-Z0-9_-]+)", url or "")
    return m.group(1) if m else (url or "").strip()


class SheetBackend:
    writable = False
    url = ""

    def read_tabs(self) -> dict[str, list[list[Any]]]:
        raise NotImplementedError

    def write_cells(self, tab: str, updates: list[tuple[int, int, Any]], *, user_entered_cols: set[int]) -> None:
        raise PlanError("此資料來源為唯讀")

    def append_rows(self, tab: str, start_row: int, rows: list[list[Any]]) -> None:
        raise PlanError("此資料來源為唯讀")

    def ensure_headers(self, tab: str, headers: list[str]) -> None:
        raise PlanError("此資料來源為唯讀")

    def write_state(self, state: dict[str, Any]) -> None:
        raise PlanError("此資料來源為唯讀")

    def write_settings(self, values: dict[str, str]) -> None:
        raise PlanError("此資料來源為唯讀")


class GoogleSheetBackend(SheetBackend):
    writable = True

    def __init__(self, spreadsheet):
        self.sh = spreadsheet
        self.url = f"https://docs.google.com/spreadsheets/d/{spreadsheet.id}/edit"
        self._ws: dict[str, Any] = {}

    def ws(self, tab: str):
        if tab not in self._ws:
            self._ws[tab] = self.sh.worksheet(tab)
        return self._ws[tab]

    def read_tabs(self) -> dict[str, list[list[Any]]]:
        import gspread

        out: dict[str, list[list[Any]]] = {}
        titles = {w.title: w for w in self.sh.worksheets()}
        self._ws.update(titles)
        ranges = [f"'{t}'" for t in (TASKS_TAB, GROUPS_TAB, PEOPLE_TAB, HOLIDAYS_TAB, SETTINGS_TAB, STATE_TAB) if t in titles]
        data = self.sh.values_batch_get(ranges, params={"valueRenderOption": "FORMATTED_VALUE"})
        for vr in data.get("valueRanges", []):
            title = vr.get("range", "").split("!")[0].strip("'")
            out[title] = vr.get("values", [])
        _ = gspread  # imported for error types in callers
        return out

    def _grow(self, ws, rows: int, cols: int) -> None:
        if rows > ws.row_count or cols > ws.col_count:
            ws.resize(rows=max(rows, ws.row_count), cols=max(cols, ws.col_count))

    def write_cells(self, tab: str, updates: list[tuple[int, int, Any]], *, user_entered_cols: set[int]) -> None:
        import gspread

        if not updates:
            return
        ws = self.ws(tab)
        self._grow(ws, max(r for r, _, _ in updates), max(c for _, c, _ in updates))
        raw, entered = [], []
        for r, c, v in updates:
            item = {"range": gspread.utils.rowcol_to_a1(r, c), "values": [[v]]}
            (entered if c in user_entered_cols else raw).append(item)
        if raw:
            ws.batch_update(raw, value_input_option="RAW")
        if entered:
            ws.batch_update(entered, value_input_option="USER_ENTERED")

    def append_rows(self, tab: str, start_row: int, rows: list[list[Any]]) -> None:
        import gspread

        if not rows:
            return
        ws = self.ws(tab)
        width = max(len(r) for r in rows)
        self._grow(ws, start_row + len(rows) - 1, width)
        rng = f"{gspread.utils.rowcol_to_a1(start_row, 1)}:{gspread.utils.rowcol_to_a1(start_row + len(rows) - 1, width)}"
        ws.update(values=rows, range_name=rng, value_input_option="USER_ENTERED")

    def ensure_headers(self, tab: str, headers: list[str]) -> None:
        ws = self.ws(tab)
        self._grow(ws, 1, len(headers))
        ws.update(values=[headers], range_name="A1", value_input_option="RAW")

    def write_state(self, state: dict[str, Any]) -> None:
        try:
            ws = self.ws(STATE_TAB)
        except Exception:  # noqa: BLE001 - worksheet missing
            ws = self.sh.add_worksheet(STATE_TAB, rows=200, cols=2)
            self._ws[STATE_TAB] = ws
            self.sh.batch_update({"requests": [{"updateSheetProperties": {
                "properties": {"sheetId": ws.id, "hidden": True}, "fields": "hidden"}}]})
        rows = [["key", "json"], ["__meta__", json.dumps(state.get("meta", {}), ensure_ascii=False)]]
        for k in sorted(state.get("items", {})):
            rows.append([k, json.dumps(state["items"][k], ensure_ascii=False, sort_keys=True)])
        ws.clear()
        self._grow(ws, len(rows), 2)
        ws.update(values=rows, range_name="A1", value_input_option="RAW")

    def write_settings(self, values: dict[str, str]) -> None:
        ws = self.ws(SETTINGS_TAB)
        existing = ws.get_all_values()
        keys = {norm_text(r[0]).lower(): i + 1 for i, r in enumerate(existing) if r}
        updates, appends = [], []
        for k, v in values.items():
            row = keys.get(k.lower())
            if row:
                updates.append({"range": f"B{row}", "values": [[v]]})
            else:
                appends.append([k, v, "（程式自動寫入，僅供顯示）"])
        if updates:
            ws.batch_update(updates, value_input_option="RAW")
        if appends:
            start = len(existing) + 1
            self._grow(ws, start + len(appends), 3)
            ws.update(values=appends, range_name=f"A{start}", value_input_option="RAW")


class XlsxBackend(SheetBackend):
    """唯讀：本機 .xlsx 檔或 Google 公開匯出（Anyone with the link can view）。"""

    def __init__(self, source: str):
        self.source = source
        if source.startswith("http"):
            self.url = source.replace("/export?format=xlsx", "/edit")

    def read_tabs(self) -> dict[str, list[list[Any]]]:
        from openpyxl import load_workbook

        if self.source.startswith("http"):
            url = sr.resolve_register_download_url(self.source)
            url = re.sub(r"([?&])gid=\d+&?", r"\1", url).rstrip("&?")
            r = requests.get(url, timeout=120)
            if not r.ok or b"<html" in r.content[:200].lower():
                raise PlanError(
                    f"無法下載 Google Sheet 公開匯出（HTTP {r.status_code}）。請確認已設為「知道連結的人可檢視」，"
                    "或設定 GOOGLE_SHEETS_JSON service account。"
                )
            wb = load_workbook(io.BytesIO(r.content), data_only=True)
        else:
            wb = load_workbook(self.source, data_only=True)
        out = {}
        for ws in wb.worksheets:
            out[ws.title] = [list(r) for r in ws.iter_rows(values_only=True)]
        return out


def open_backend(plan_cfg: dict[str, Any], xlsx_override: str = "") -> SheetBackend:
    if xlsx_override:
        return XlsxBackend(xlsx_override)
    sheet_cfg = plan_cfg.get("sheet") or {}
    gc = gspread_client()
    sid = sheet_id_from_url(str(sheet_cfg.get("sheet_id") or sheet_cfg.get("url") or ""))
    title = str(sheet_cfg.get("title") or "C27_Project_Plan")
    if gc is not None:
        sh = gc.open_by_key(sid) if sid else gc.open(title)
        return GoogleSheetBackend(sh)
    export_url = str(sheet_cfg.get("public_export_url") or "")
    if not export_url and sid:
        export_url = f"https://docs.google.com/spreadsheets/d/{sid}/export?format=xlsx"
    if not export_url:
        raise PlanError("沒有 Google 憑證（GOOGLE_SHEETS_JSON），也沒有設定 plan.sheet.sheet_id / public_export_url")
    print("注意：沒有 Google service account 憑證，改用公開匯出唯讀模式（不會寫回 Sheet）。")
    return XlsxBackend(export_url)


# ---------------------------------------------------------------------------
# Jira side
# ---------------------------------------------------------------------------

@dataclass
class JiraLink:
    link_id: str
    other_key: str
    other_is_blocker: bool  # True = other blocks this issue


@dataclass
class JiraIssueView:
    key: str
    summary: str
    plan_id: str
    name: str
    status: str
    labels: list[str]
    parent_key: str
    assignee_id: str
    assignee_name: str
    start: str
    due: str
    desc_text: str
    links: list[JiraLink]
    is_epic: bool
    reporter: str = ""
    desc_full: str = ""

    @property
    def removed(self) -> bool:
        return "plan-removed" in self.labels


def adf_to_text(node: Any, *, stop_at_marker: bool = True) -> str:
    """ADF → 純文字（段落換行）；預設遇到同步標記段落即停止（只取使用者描述）。"""
    lines: list[str] = []

    def inline_text(n: Any) -> str:
        if not isinstance(n, dict):
            return ""
        if n.get("type") == "text":
            return str(n.get("text") or "")
        if n.get("type") == "hardBreak":
            return "\n"
        return "".join(inline_text(c) for c in n.get("content") or [])

    def walk(blocks: list[Any]) -> bool:
        for b in blocks or []:
            if not isinstance(b, dict):
                continue
            t = b.get("type")
            if t in ("paragraph", "heading"):
                text = inline_text(b)
                if stop_at_marker and text.strip().startswith(DESC_MARKER):
                    return True
                lines.extend(text.split("\n"))
            elif t == "rule":
                continue
            elif t in ("bulletList", "orderedList", "listItem", "blockquote", "panel", "table",
                       "tableRow", "tableCell", "tableHeader", "expand"):
                if walk(b.get("content") or []):
                    return True
            elif t == "codeBlock":
                lines.extend(inline_text(b).split("\n"))
        return False

    if isinstance(node, dict):
        walk(node.get("content") or [])
    elif isinstance(node, str):
        return norm_multiline(node.split(DESC_MARKER)[0] if stop_at_marker else node)
    return norm_multiline("\n".join(lines))


def build_issue_description(task: PlanTask, group: Group | None, sheet_url: str) -> dict[str, Any]:
    content: list[dict[str, Any]] = []
    for line in (task.description or "").split("\n"):
        content.append({"type": "paragraph", "content": [{"type": "text", "text": line}] if line else []})
    content.append({"type": "rule"})
    meta = [
        f"{DESC_MARKER} 以下由 Google Sheet 自動產生，請勿在此段落編輯",
        f"Plan ID: {task.id}",
        f"Group: {task.group}" + (f" ({group.epic_name})" if group else ""),
        f"Type: {task.type}    Duration: {task.duration} 工作天" + ("（估計值，待確認）" if task.estimate else ""),
        f"Predecessors: {format_predecessors(task.preds) or '-'}",
    ]
    if sheet_url:
        meta.append(f"Sheet: {sheet_url}")
    for line in meta:
        content.append({"type": "paragraph", "content": [{"type": "text", "text": line}]})
    return {"type": "doc", "version": 1, "content": content}


def issue_view_from_json(issue: dict[str, Any], start_field: str, epic_type_names: set[str]) -> JiraIssueView:
    f = issue.get("fields") or {}
    summary = str(f.get("summary") or "")
    m = PLAN_SUMMARY_RE.match(summary.strip())
    plan_id, name = (m.group(1), m.group(2).strip()) if m else ("", summary.strip())
    links: list[JiraLink] = []
    for ln in f.get("issuelinks") or []:
        if (ln.get("type") or {}).get("name") != "Blocks":
            continue
        if ln.get("inwardIssue"):
            links.append(JiraLink(str(ln.get("id")), ln["inwardIssue"]["key"], True))
        elif ln.get("outwardIssue"):
            links.append(JiraLink(str(ln.get("id")), ln["outwardIssue"]["key"], False))
    itype = f.get("issuetype") or {}
    assignee = f.get("assignee") or {}
    return JiraIssueView(
        key=issue["key"],
        summary=summary,
        plan_id=plan_id,
        name=name,
        status=normalize_jira_status((f.get("status") or {}).get("name", "")),
        labels=list(f.get("labels") or []),
        parent_key=(f.get("parent") or {}).get("key", ""),
        assignee_id=assignee.get("accountId", "") or "",
        assignee_name=assignee.get("displayName", "") or "",
        start=str(f.get(start_field) or "")[:10],
        due=str(f.get("duedate") or "")[:10],
        desc_text=adf_to_text(f.get("description")),
        links=links,
        is_epic=(itype.get("hierarchyLevel") == 1) or (itype.get("name") in epic_type_names),
        reporter=(f.get("reporter") or {}).get("displayName", "") or "",
        desc_full=adf_to_text(f.get("description"), stop_at_marker=False),
    )


def looks_like_register_issue(view: JiraIssueView) -> bool:
    """Register 單（summary 前綴符合 ^[A-Z][A-Z0-9]*-\\d+$）絕不由計畫同步處理。"""
    return sr.parse_register_id_from_summary(view.summary) is not None and not view.plan_id


# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

DEFAULT_PLAN_CFG: dict[str, Any] = {
    "sheet": {"sheet_id": "", "title": "C27_Project_Plan", "share_with": ["bob.lin@qsitw.com"], "public_export_url": ""},
    "confluence": {
        "space_key": "projectmod",
        "space_id": "1409028",
        "parent_page_id": "14057511",
        "page_title": "C27 Project Plan",
        "page_id": "",
        "show_discussion": True,
        "version_message": "Synced from C27 plan sheet",
    },
    "jira": {
        "project_key": "PMWC",
        "epic_issue_type": "大型工作",
        "task_issue_type": "任務",
        "start_date_field": "customfield_10015",
        "due_date_field": "duedate",
        "label": "c27-plan",
        "default_assignee_account_id": "712020:a50a0ee0-1928-47a7-a135-a07c21c233ce",
        "on_removed": "close",
        "done_status": JIRA_DONE,
        "max_removals": 10,
    },
    "notify": {},
}


def deep_merge(base: dict[str, Any], over: dict[str, Any]) -> dict[str, Any]:
    out = copy.deepcopy(base)
    for k, v in (over or {}).items():
        if isinstance(v, dict) and isinstance(out.get(k), dict):
            out[k] = deep_merge(out[k], v)
        else:
            out[k] = copy.deepcopy(v)
    return out


def load_plan_config(config_path: Path, plan_config_path: Path | None, *, need_token: bool = True) -> dict[str, Any]:
    """atlassian / notify 取自 config.yaml；計畫設定只取 `plan:` 區塊（避免誤用 Register 的 page_id）。"""
    sr.load_dotenv(SCRIPT_DIR / ".env")
    base: dict[str, Any] = {}
    if config_path and config_path.exists():
        base = yaml.safe_load(config_path.read_text(encoding="utf-8")) or {}
    plan_section = base.get("plan") or {}
    if plan_config_path and plan_config_path.exists():
        extra = yaml.safe_load(plan_config_path.read_text(encoding="utf-8")) or {}
        plan_section = deep_merge(plan_section, extra.get("plan", extra))
        for k in ("atlassian", "silent_mode", "notify"):
            if k in extra and k not in base:
                base[k] = extra[k]
    plan = deep_merge(DEFAULT_PLAN_CFG, plan_section)
    page_id = str(plan["confluence"].get("page_id") or "")
    if page_id in REGISTER_PAGE_IDS or str(plan["confluence"].get("page_title", "")).strip() in PROTECTED_PAGE_TITLES:
        raise PlanError("plan.confluence 指向 Register 的 mail_checking 頁面，已拒絕（避免覆寫）")
    if plan["jira"].get("on_removed") not in ("close", "delete"):
        raise PlanError("plan.jira.on_removed 只能是 close 或 delete")
    cfg = {
        "atlassian": base.get("atlassian") or {
            "site_url": "https://qsiaiot.atlassian.net",
            "cloud_id": "2657b5c5-bbc7-4469-9e06-6bc58db3de73",
            "email": "bob.lin@qsitw.com",
        },
        "silent_mode": base.get("silent_mode", True),
        "notify": deep_merge(base.get("notify") or {}, plan.get("notify") or {}),
        "plan": plan,
    }
    token = os.environ.get("ATLASSIAN_API_TOKEN", "").strip()
    if need_token and not token:
        raise PlanError("缺少 ATLASSIAN_API_TOKEN（.env 或環境變數）")
    cfg["_api_token"] = token
    return cfg


class PlanAtlassianClient(sr.AtlassianClient):
    def confluence_post(self, path: str, payload: dict) -> Any:
        r = self.session.post(f"{self.confluence_v2}{path}", json=payload, timeout=60)
        if not r.ok:
            raise RuntimeError(f"Confluence POST failed: {r.status_code} {r.text}")
        return r.json()


# ---------------------------------------------------------------------------
# Three-way merge
# ---------------------------------------------------------------------------

def merge_scalar(field_name: str, a: str, b: str, s: str | None, resolution: str = "") -> str:
    """回傳 same / push（Sheet→Jira）/ pull（Jira→Sheet）/ conflict。"""
    if a == b:
        return "same"
    if resolution == "sheet":
        return "push"
    if resolution == "jira":
        return "pull"
    if s is None:
        return "push" if FIELD_MASTER.get(field_name, "sheet") == "sheet" else "pull"
    if a == s:
        return "pull"
    if b == s:
        return "push"
    return "conflict"


def merge_sets(a: set[str], b: set[str], s: set[str] | None) -> tuple[set[str], set[str], set[str]]:
    """前置任務集合三方合併。回傳 (合併結果, Jira 需新增到 Sheet 的, Jira 需從 Sheet 移除的)。"""
    if s is None:
        result = a | b
    else:
        result = (s | (a - s) | (b - s)) - (s - a) - (s - b)
    return result, result - a, a - result


@dataclass
class FieldDecision:
    field: str
    decision: str
    sheet_value: str
    jira_value: str
    snap_value: str | None


# ---------------------------------------------------------------------------
# Report
# ---------------------------------------------------------------------------

@dataclass
class PlanReport:
    direction: str = "both"
    dry_run: bool = False
    created_epics: list[str] = field(default_factory=list)
    created_issues: list[str] = field(default_factory=list)
    updated_issues: list[str] = field(default_factory=list)
    closed_issues: list[str] = field(default_factory=list)
    deleted_issues: list[str] = field(default_factory=list)
    links_added: list[str] = field(default_factory=list)
    links_removed: list[str] = field(default_factory=list)
    pulled: list[str] = field(default_factory=list)
    new_from_jira: list[str] = field(default_factory=list)
    conflicts: list[str] = field(default_factory=list)
    pending: list[str] = field(default_factory=list)
    schedule_changes: list[str] = field(default_factory=list)
    assigned_ids: list[str] = field(default_factory=list)
    sheet_writes: int = 0
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    summary_lines: list[str] = field(default_factory=list)
    confluence_url: str = ""
    confluence_action: str = ""
    sheet_url: str = ""

    def has_changes(self) -> bool:
        return any([
            self.created_epics, self.created_issues, self.updated_issues, self.closed_issues,
            self.deleted_issues, self.links_added, self.links_removed, self.pulled,
            self.new_from_jira, self.schedule_changes, self.assigned_ids,
        ])

    def render(self) -> str:
        def sec(title: str, items: list[str], limit: int = 60) -> list[str]:
            if not items:
                return []
            out = [f"■ {title}（{len(items)}）"]
            out += [f"  - {x}" for x in items[:limit]]
            if len(items) > limit:
                out.append(f"  …另 {len(items) - limit} 筆")
            return out

        lines = [
            f"C27 Project Plan 同步結果（direction={self.direction}{'，DRY-RUN 未寫入' if self.dry_run else ''}）",
            f"時間：{datetime.now(TW_TZ):%Y-%m-%d %H:%M} (UTC+8)",
        ]
        if self.sheet_url:
            lines.append(f"Google Sheet：{self.sheet_url}")
        if self.confluence_url:
            lines.append(f"Confluence：{self.confluence_url}（{self.confluence_action or '-'}）")
        lines.append("")
        lines += self.summary_lines
        lines.append("")
        lines += sec("衝突（兩邊都改，未覆寫，請在 Sync Conflict 欄填 use sheet / use jira）", self.conflicts)
        lines += sec("Jira → Sheet 拉回", self.pulled)
        lines += sec("Jira 新建、已加入 Sheet 的列", self.new_from_jira)
        lines += sec("新建 Epic", self.created_epics)
        lines += sec("新建任務", self.created_issues)
        lines += sec("更新任務", self.updated_issues)
        lines += sec("關閉（Sheet 已刪除）", self.closed_issues)
        lines += sec("刪除（Sheet 已刪除）", self.deleted_issues)
        lines += sec("新增相依（Blocks）", self.links_added)
        lines += sec("移除相依", self.links_removed)
        lines += sec("排程變更", self.schedule_changes)
        lines += sec("自動配發 ID", self.assigned_ids)
        lines += sec("待處理（本次方向未同步）", self.pending)
        lines += sec("警告", self.warnings)
        lines += sec("錯誤", self.errors)
        if not self.has_changes() and not self.conflicts:
            lines.append("無變更。")
        return "\n".join(lines)


# ---------------------------------------------------------------------------
# Sync engine
# ---------------------------------------------------------------------------

class PlanSync:
    def __init__(self, cfg: dict[str, Any], backend: SheetBackend, client: Any, *,
                 direction: str = "both", dry_run: bool = False, allow_mass_removal: bool = False,
                 state_file: Path | None = None):
        self.cfg = cfg
        self.pcfg = cfg["plan"]
        self.jcfg = self.pcfg["jira"]
        self.backend = backend
        self.client = client
        self.direction = direction
        self.dry_run = dry_run
        self.allow_mass_removal = allow_mass_removal
        self.state_file = state_file or (SCRIPT_DIR / ".plan_cache" / "sync_state.json")
        self.report = PlanReport(direction=direction, dry_run=dry_run, sheet_url=backend.url)
        self.label = str(self.jcfg.get("label") or "c27-plan")
        self.site_url = str(cfg["atlassian"]["site_url"]).rstrip("/")
        self.start_field = str(self.jcfg.get("start_date_field") or "customfield_10015")
        self.epic_type = sr.resolve_issuetype_for_api(self.jcfg.get("epic_issue_type", "大型工作"), role="epic")
        self.task_type = sr.resolve_issuetype_for_api(self.jcfg.get("task_issue_type", "任務"), role="task")
        self.link_flip = False
        self.link_verified = False
        self.sheet_updates: dict[tuple[int, str], Any] = {}
        self._user_cache: dict[str, str] = {}

    # ----- helpers -----
    @property
    def do_pull(self) -> bool:
        return self.direction in ("pull", "both")

    @property
    def do_push(self) -> bool:
        return self.direction in ("push", "both")

    def browse(self, key: str) -> str:
        return f"{self.site_url}/browse/{key}"

    def set_cell(self, task: PlanTask, col: str, value: Any) -> None:
        current = task.cells.get(col, "")
        if norm_text(value) == norm_text(current):
            return
        task.cells[col] = norm_text(value)
        if task.row:
            self.sheet_updates[(task.row, col)] = value

    # ----- people -----
    def resolve_owner(self, owner: str, people: list[Person]) -> str | None:
        """Owner 名稱 → accountId；空白 → 預設受託人；找不到 → None。"""
        default = str(self.jcfg.get("default_assignee_account_id") or "")
        name = (owner or "").strip()
        if not name:
            return default
        low = name.lower()
        for p in people:
            if low in (p.name.lower(), p.email.lower()) or (p.email and low == p.email.split("@")[0].lower()):
                if p.account_id:
                    return p.account_id
                if p.email:
                    return self.lookup_account(p.email)
        if re.match(r"^[0-9a-f]{6,}:[0-9a-f-]{20,}$", name):
            return name
        return self.lookup_account(name)

    def lookup_account(self, query: str) -> str | None:
        if query in self._user_cache:
            return self._user_cache[query] or None
        acc = ""
        try:
            users = self.client.jira_get("/user/search", query=query) or []
            active = [u for u in users if u.get("active", True) and u.get("accountType") != "app"]
            if len(active) == 1:
                acc = active[0].get("accountId", "")
        except Exception as exc:  # noqa: BLE001
            self.report.warnings.append(f"查詢 Jira 使用者 {query!r} 失敗：{exc}")
        self._user_cache[query] = acc
        return acc or None

    def owner_display(self, account_id: str, fallback: str, people: list[Person]) -> str:
        for p in people:
            if p.account_id and p.account_id == account_id:
                return p.name
        return fallback

    # ----- Jira fetch -----
    def fetch_plan_issues(self) -> tuple[dict[str, JiraIssueView], dict[str, JiraIssueView]]:
        project = self.jcfg["project_key"]
        fields = ["summary", "status", "labels", "parent", "issuetype", "assignee", self.start_field,
                  "duedate", "description", "issuelinks", "reporter"]
        epic_names = set(sr.jql_issuetype_name_candidates(self.epic_type))
        issues = sr._paginate_jira_search(
            self.client, f'project = {project} AND labels = {sr._jql_quote(self.label)}', fields=fields)
        views = [issue_view_from_json(i, self.start_field, epic_names) for i in issues]
        epics = {v.key: v for v in views if v.is_epic}
        tasks = {v.key: v for v in views if not v.is_epic}
        if epics:
            keys = ", ".join(sorted(epics))
            extra = sr._paginate_jira_search(
                self.client,
                f"project = {project} AND parent in ({keys}) AND (labels IS EMPTY OR labels != {sr._jql_quote(self.label)})",
                fields=fields,
            )
            for i in extra:
                v = issue_view_from_json(i, self.start_field, epic_names)
                if not v.is_epic:
                    tasks.setdefault(v.key, v)
        for key in list(tasks):
            if looks_like_register_issue(tasks[key]):
                self.report.warnings.append(f"{key} 看起來是 Register 單（{tasks[key].summary[:40]}），計畫同步略過")
                del tasks[key]
        return epics, tasks

    # ----- state -----
    def load_state(self, data: PlanData) -> dict[str, Any]:
        if data.state.get("items") or data.state.get("meta"):
            return data.state
        if self.state_file.exists():
            try:
                return json.loads(self.state_file.read_text(encoding="utf-8"))
            except ValueError:
                pass
        return {"meta": {}, "items": {}}

    def save_state(self, state: dict[str, Any]) -> None:
        if self.dry_run:
            return
        if self.backend.writable:
            self.backend.write_state(state)
        self.state_file.parent.mkdir(parents=True, exist_ok=True)
        self.state_file.write_text(json.dumps(state, ensure_ascii=False, indent=1, sort_keys=True), encoding="utf-8")

    # ----- values -----
    def sheet_values(self, t: PlanTask, people: list[Person], synced_ids: set[str]) -> dict[str, str]:
        owner = self.resolve_owner(t.owner, people)
        return {
            "name": t.name.strip(),
            "description": norm_multiline(t.description),
            "priority": t.priority,
            "group": t.group,
            "status": SHEET_TO_JIRA_STATUS.get(t.status, JIRA_TODO),
            "owner": owner if owner is not None else "?",
            "start": t.start.isoformat() if t.start else "",
            "finish": t.finish.isoformat() if t.finish else "",
            "preds": ",".join(sorted(p.id for p in t.preds if p.id in synced_ids)),
        }

    def jira_values(self, v: JiraIssueView, epic_code_by_key: dict[str, str], key_to_id: dict[str, str]) -> dict[str, str]:
        prio = ""
        for lab in v.labels:
            if lab.lower().startswith("prio-"):
                prio = normalize_priority(lab[5:])
        preds = sorted(key_to_id[l.other_key] for l in v.links if l.other_is_blocker and l.other_key in key_to_id)
        return {
            "name": v.name.strip(),
            "description": v.desc_text,
            "priority": prio,
            "group": epic_code_by_key.get(v.parent_key, ""),
            "status": v.status,
            "owner": v.assignee_id,
            "start": v.start,
            "finish": v.due,
            "preds": ",".join(preds),
        }

    # ----- main -----
    def run(self) -> PlanReport:
        rep = self.report
        tabs = self.backend.read_tabs()
        data = build_plan_data(tabs)
        if not data.tasks:
            raise PlanError("Tasks 分頁沒有任何任務，為安全起見中止（不會關閉任何 Jira 單）")
        missing_groups = sorted({t.group for t in data.tasks if t.group and t.group not in data.groups})
        for g in missing_groups:
            data.groups[g] = Group(g, f"PLAN-{g}")
            rep.warnings.append(f"Group {g} 不在 Groups 分頁，Epic 名稱暫用 PLAN-{g}")
        for t in assign_missing_ids(data.tasks):
            rep.assigned_ids.append(f"列 {t.row}: {t.id} {t.name}")
            self.set_cell(t, "id", t.id)
        validate_ids(data.tasks)
        resolve_predecessors(data.tasks)
        cal = project_calendar(data)
        rep.warnings += compute_schedule(data.tasks, cal)
        state = self.load_state(data)
        snap_items: dict[str, dict[str, Any]] = state.setdefault("items", {})

        epics, jira_tasks = self.fetch_plan_issues()
        epic_by_code: dict[str, JiraIssueView] = {}
        for e in epics.values():
            for lab in e.labels:
                if lab.lower().startswith("plan-epic-"):
                    epic_by_code[lab[10:].upper()] = e
        epic_code_by_key = {e.key: code for code, e in epic_by_code.items()}

        # ----- map sheet rows <-> Jira issues -----
        by_id = data.task_by_id()
        by_plan_id = {v.plan_id: v for v in jira_tasks.values() if v.plan_id and not v.removed}
        pair: dict[str, JiraIssueView] = {}
        for t in data.tasks:
            v = jira_tasks.get(t.jira_key) if t.jira_key else None
            if v is None and t.jira_key and t.jira_key not in jira_tasks:
                rep.warnings.append(f"{t.id}: Jira Key {t.jira_key} 不在計畫單中（已刪除或缺 label），將重新比對")
            if v is None:
                v = by_plan_id.get(t.id)
            if v is not None:
                pair[t.id] = v
                if t.jira_key != v.key:
                    t.jira_key = v.key
                    self.set_cell(t, "jira_key", v.key)
        mapped_keys = {v.key for v in pair.values()}

        # ----- issues in Jira but not in sheet -----
        removed_candidates: list[tuple[str, JiraIssueView | None]] = []
        for key, v in jira_tasks.items():
            if key in mapped_keys or v.removed:
                continue
            if key in snap_items:
                removed_candidates.append((key, v))
            elif self.do_pull:
                self.adopt_jira_issue(data, v, epic_code_by_key, rep)
            else:
                rep.pending.append(f"{key} {v.summary[:50]}：Jira 新建，需執行 pull 才會加入 Sheet")
        for key in list(snap_items):
            if key not in jira_tasks and key not in mapped_keys and not any(t.jira_key == key for t in data.tasks):
                snap_items.pop(key, None)  # deleted in Jira manually; forget
        by_id = data.task_by_id()
        for t in data.tasks:
            if t.jira_key and t.jira_key in jira_tasks and t.id not in pair:
                pair[t.id] = jira_tasks[t.jira_key]

        synced_ids = {t.id for t in data.tasks if t.sync}
        key_to_id = {v.key: tid for tid, v in pair.items()}

        # ----- pull phase: three-way merge per field -----
        decisions: dict[str, dict[str, FieldDecision]] = {}
        pred_merge: dict[str, tuple[set[str], set[str], set[str]]] = {}
        for t in data.tasks:
            v = pair.get(t.id)
            if not t.sync or v is None:
                continue
            snap = snap_items.get(v.key)
            snap_fields = (snap or {}).get("fields")
            a = self.sheet_values(t, data.people, synced_ids)
            b = self.jira_values(v, epic_code_by_key, key_to_id)
            res = t.resolution()
            row_dec: dict[str, FieldDecision] = {}
            for f in SYNC_FIELDS:
                s = snap_fields.get(f) if snap_fields is not None else None
                if f == "preds":
                    A = set(filter(None, a[f].split(",")))
                    B = set(filter(None, b[f].split(",")))
                    S = set(filter(None, s.split(","))) if s is not None else None
                    result, add_sheet, rm_sheet = merge_sets(A, B, S)
                    pred_merge[t.id] = (result, add_sheet, rm_sheet)
                    if A == B:
                        dec = "same"
                    elif not add_sheet and not rm_sheet:
                        dec = "push"
                    elif result == B:
                        dec = "pull"
                    else:
                        dec = "merge"
                    row_dec[f] = FieldDecision(f, dec, a[f], b[f], s)
                    continue
                if f == "owner" and a[f] == "?":
                    rep.warnings.append(f"{t.id}: Owner {t.owner!r} 找不到對應的 Jira 帳號（請在 People 分頁補 Jira Account ID）")
                    row_dec[f] = FieldDecision(f, "skip", a[f], b[f], s)
                    continue
                dec = merge_scalar(f, a[f], b[f], s, res)
                row_dec[f] = FieldDecision(f, dec, a[f], b[f], s)
            decisions[t.id] = row_dec

        if self.do_pull:
            for t in data.tasks:
                if t.id in decisions:
                    self.apply_pulls(t, decisions[t.id], pred_merge.get(t.id), pair[t.id], data, cal)
            resolve_predecessors(data.tasks)
            rep.warnings += compute_schedule(data.tasks, cal)

        # ----- schedule summary & validation -----
        self.summarize_schedule(data, cal)

        # ----- push phase -----
        final_jira: dict[str, dict[str, str]] = {}
        if self.do_push:
            final_jira = self.push(data, pair, decisions, epic_by_code, synced_ids)
            self.handle_removed(removed_candidates, data, snap_items)
        else:
            for key, v in removed_candidates:
                rep.pending.append(f"{key} {v.summary[:50] if v else ''}：Sheet 已刪除，需執行 push 才會{self.jcfg['on_removed']}")

        # ----- conflicts column, computed columns -----
        now = datetime.now(TW_TZ).strftime("%Y-%m-%d %H:%M")
        key_to_id = {v.key: tid for tid, v in pair.items()}
        for t in data.tasks:
            self.set_cell(t, "start", t.start.isoformat() if t.start else "")
            self.set_cell(t, "finish", t.finish.isoformat() if t.finish else "")
            row_dec = decisions.get(t.id, {})
            conflicts = [d for d in row_dec.values() if d.decision == "conflict"]
            notes = []
            for d in conflicts:
                notes.append(f"{FIELD_LABELS[d.field]}: Sheet={d.sheet_value[:40] or '(空)'} / Jira={d.jira_value[:40] or '(空)'}")
                rep.conflicts.append(f"{t.id} {FIELD_LABELS[d.field]}: Sheet={d.sheet_value[:40]!r} Jira={d.jira_value[:40]!r}")
            if t.errors:
                notes.append("⚠ " + "；".join(t.errors))
                for e in t.errors:
                    rep.warnings.append(f"{t.id}: {e}")
            conflict_text = ""
            if notes:
                conflict_text = " | ".join(notes)
                if conflicts:
                    conflict_text += " → 在此欄輸入 use sheet 或 use jira 決定以哪邊為準"
            if self.backend.writable or self.dry_run:
                self.set_cell(t, "conflict", conflict_text)
            if t.id in decisions and (self.do_push or self.do_pull) and not self.dry_run:
                changed = [d.field for d in decisions[t.id].values() if d.decision in ("push", "pull", "merge")]
                if changed or t.row == 0:
                    self.set_cell(t, "last_sync", now)

        # ----- new snapshot -----
        if not self.dry_run:
            synced_ids = {t.id for t in data.tasks if t.sync}
            key_to_id = {v.key: tid for tid, v in pair.items()}
            for t in data.tasks:
                if not t.sync or not t.jira_key:
                    continue
                a = self.sheet_values(t, data.people, synced_ids)
                b = final_jira.get(t.id)
                old = (snap_items.get(t.jira_key) or {}).get("fields") or {}
                new_fields = dict(old)
                row_dec = decisions.get(t.id, {})
                for f in SYNC_FIELDS:
                    dec = row_dec.get(f)
                    if b is not None and a[f] == b.get(f):
                        new_fields[f] = a[f]
                    elif dec and dec.decision == "pull" and self.do_pull:
                        new_fields[f] = dec.jira_value
                    elif dec and dec.decision == "push" and self.do_push and b is not None:
                        new_fields[f] = b.get(f, a[f])
                snap_items[t.jira_key] = {"id": t.id, "fields": new_fields}
            state["meta"] = {**state.get("meta", {}), "last_sync": now, "direction": self.direction}

        # ----- write sheet -----
        self.write_sheet(data)

        # ----- Confluence -----
        if self.do_push:
            try:
                self.update_confluence(data, cal, state)
            except Exception as exc:  # noqa: BLE001
                rep.errors.append(f"Confluence 更新失敗：{exc}")
        self.save_state(state)
        return rep

    # ----- adopt Jira-created issue -----
    def adopt_jira_issue(self, data: PlanData, v: JiraIssueView, epic_code_by_key: dict[str, str], rep: PlanReport) -> None:
        if not self.backend.writable and not self.dry_run:
            rep.pending.append(f"{v.key} {v.summary[:50]}：Jira 新建，但 Sheet 唯讀無法加入")
            return
        group = epic_code_by_key.get(v.parent_key, "") or "JR"
        existing = {t.id for t in data.tasks}
        tid = v.plan_id if v.plan_id and v.plan_id not in existing else next_plan_id(group, existing)
        dur = 1
        if v.start and v.due:
            cal = project_calendar(data)
            dur = max(0, cal.workdays_between(date.fromisoformat(v.start), date.fromisoformat(v.due)))
        prio = ""
        for lab in v.labels:
            if lab.lower().startswith("prio-"):
                prio = normalize_priority(lab[5:])
        status = {JIRA_TODO: "Not Started", JIRA_DOING: "In Progress", JIRA_DONE: "Done"}.get(v.status, "Not Started")
        cells = {
            "id": tid, "group": group, "name": v.name, "type": TYPE_MILESTONE if dur == 0 else TYPE_TASK,
            "priority": prio, "owner": self.owner_display(v.assignee_id, v.assignee_name, data.people),
            "duration": str(dur), "preds": "", "constraint": v.start or "", "status": status,
            "sync": "Y", "jira_key": v.key, "description": v.desc_text,
            "proposed_by": f"Jira: {v.reporter}".strip(), "estimate": "N",
            "notes": f"由 Jira {v.key} 新增", "conflict": "", "last_sync": "",
        }
        t = PlanTask(row=0, id=tid, cells=cells, group=group, name=v.name, type=cells["type"], priority=prio,
                     owner=cells["owner"], duration=dur, constraint_text=cells["constraint"], status=status,
                     sync=True, jira_key=v.key, description=v.desc_text)
        data.tasks.append(t)
        rep.new_from_jira.append(f"{v.key} → {tid} {v.name[:50]}")

    # ----- apply Jira->Sheet changes to the model -----
    def apply_pulls(self, t: PlanTask, row_dec: dict[str, FieldDecision], pm, v: JiraIssueView,
                    data: PlanData, cal: WorkCalendar) -> None:
        pulled = []
        for f, d in row_dec.items():
            if d.decision != "pull" or f in ("start", "finish", "preds"):
                continue
            if f == "name":
                t.name = d.jira_value
                self.set_cell(t, "name", t.name)
            elif f == "description":
                t.description = d.jira_value
                self.set_cell(t, "description", t.description)
            elif f == "priority" and d.jira_value:
                t.priority = d.jira_value
                self.set_cell(t, "priority", t.priority)
            elif f == "group" and d.jira_value:
                t.group = d.jira_value
                self.set_cell(t, "group", t.group)
            elif f == "status":
                new = {JIRA_TODO: "Not Started", JIRA_DOING: "In Progress", JIRA_DONE: "Done"}.get(d.jira_value)
                if not new:
                    continue
                t.status = new
                self.set_cell(t, "status", new)
                if new == "Done":
                    self.set_cell(t, "pct", "100%")
            elif f == "owner":
                t.owner = self.owner_display(d.jira_value, v.assignee_name, data.people) if d.jira_value else ""
                self.set_cell(t, "owner", t.owner)
            pulled.append(f"{FIELD_LABELS[f]}={d.jira_value[:30]!r}")
        ds, df = row_dec.get("start"), row_dec.get("finish")
        if (ds and ds.decision == "pull") or (df and df.decision == "pull"):
            js = date.fromisoformat(v.start) if v.start else t.start
            jd = date.fromisoformat(v.due) if v.due else t.finish
            if ds and ds.decision == "pull" and js:
                t.constraint_text = js.isoformat()
                self.set_cell(t, "constraint", t.constraint_text)
                pulled.append(f"Start Constraint={js}")
            if js and jd:
                new_dur = max(0, cal.workdays_between(js, jd))
                if t.is_milestone and js == jd:
                    new_dur = 0
                if new_dur != t.duration:
                    pulled.append(f"Duration {t.duration}→{new_dur}d")
                    t.duration = new_dur
                    self.set_cell(t, "duration", new_dur)
        if pm is not None:
            result, add_sheet, rm_sheet = pm
            if add_sheet or rm_sheet:
                keep = [p for p in t.preds if p.id not in rm_sheet]
                keep += [Pred(i) for i in sorted(add_sheet)]
                # keep draft (unsynced) predecessors that are not part of the compare set
                t.preds_text = format_predecessors(keep)
                self.set_cell(t, "preds", t.preds_text)
                if add_sheet:
                    pulled.append("Predecessors +" + ",".join(sorted(add_sheet)))
                if rm_sheet:
                    pulled.append("Predecessors -" + ",".join(sorted(rm_sheet)))
        if pulled:
            self.report.pulled.append(f"{t.id} ({v.key}): " + "; ".join(pulled))

    # ----- schedule summary -----
    def summarize_schedule(self, data: PlanData, cal: WorkCalendar) -> None:
        rep = self.report
        for t in data.tasks:
            prev_s, prev_f = t.cells.get("start", ""), t.cells.get("finish", "")
            new_s = t.start.isoformat() if t.start else ""
            new_f = t.finish.isoformat() if t.finish else ""
            if (prev_s or prev_f) and (prev_s != new_s or prev_f != new_f):
                rep.schedule_changes.append(f"{t.id}: {prev_s}~{prev_f} → {new_s}~{new_f}")
        starts = [t.start for t in data.tasks if t.start]
        finishes = [t.finish for t in data.tasks if t.finish]
        t0 = next((t for t in data.tasks if t.id == "PLN-G-05"), None)
        lines = [
            f"Project Start：{cal.base}" + ("（佔位日期，請在 Settings 分頁設定）" if "placeholder" in data.settings.get("project start note", "").lower() else ""),
            f"整體：{min(starts) if starts else '-'} ~ {max(finishes) if finishes else '-'}；任務 {len(data.tasks)} 筆"
            f"（同步 {sum(1 for t in data.tasks if t.sync)}、提案/討論中 {sum(1 for t in data.tasks if not t.sync)}）",
        ]
        if t0 and t0.finish:
            lines.append(f"T0（前置條件就緒）：{t0.finish}")
        spans = group_spans(data.tasks)
        for code, g in data.groups.items():
            if code in spans:
                s, f = spans[code]
                dur = sum(t.duration for t in data.tasks if t.group == code and t.type == TYPE_TASK)
                lines.append(f"  {code:<4} {g.epic_name}：{s} ~ {f}（任務工期合計 {dur}d）")
        for code, expected, actual, ok in validate_group_totals(data.tasks, data.groups):
            if ok:
                lines.append(f"  ✓ {code} 工期合計 {actual}d = 承諾 {expected}d（{expected / 5:g}W）")
            else:
                msg = f"{code} 工期合計 {actual}d ≠ 承諾 {expected}d（{expected / 5:g}W），請確認"
                lines.append(f"  ⚠ {msg}")
                rep.warnings.append(msg)
        rep.summary_lines = lines

    # ----- push -----
    def desired_labels(self, t: PlanTask, existing: list[str]) -> list[str]:
        keep = [l for l in existing if not (l == self.label or l.startswith("plan-") or l.startswith("prio-"))]
        managed = [self.label, f"plan-{t.group.lower()}"]
        if t.priority:
            managed.append(f"prio-{t.priority.lower()}")
        if t.type == TYPE_MILESTONE:
            managed.append("plan-milestone")
        if t.type == TYPE_PROCUREMENT:
            managed.append("plan-procurement")
        if t.estimate:
            managed.append("plan-estimate")
        return sorted(set(keep + managed))

    def ensure_epics(self, data: PlanData, epic_by_code: dict[str, JiraIssueView], synced: list[PlanTask]) -> dict[str, str]:
        out: dict[str, str] = {}
        project = self.jcfg["project_key"]
        for code in sorted({t.group for t in synced}):
            g = data.groups.get(code) or Group(code, f"PLAN-{code}")
            children = [t for t in synced if t.group == code and t.start and t.finish]
            start = min(t.start for t in children).isoformat() if children else None
            due = max(t.finish for t in children).isoformat() if children else None
            existing = epic_by_code.get(code)
            labels = [self.label, f"plan-epic-{code.lower()}"]
            fields: dict[str, Any] = {}
            if existing is None:
                fields = {
                    "project": {"key": project}, "issuetype": {"name": self.epic_type},
                    "summary": g.epic_name, "labels": labels,
                    "description": sr.text_to_adf(g.description or g.epic_name),
                }
                if start:
                    fields[self.start_field] = start
                    fields["duedate"] = due
                if self.dry_run:
                    out[code] = f"NEW-EPIC-{code}"
                    self.report.created_epics.append(f"{g.epic_name} (dry-run)")
                    continue
                created = self.client.jira_post("/issue", {"fields": fields})
                out[code] = created["key"]
                self.report.created_epics.append(f"{created['key']} {g.epic_name}")
                continue
            out[code] = existing.key
            if existing.summary != g.epic_name:
                fields["summary"] = g.epic_name
            if start and (existing.start != start or existing.due != due):
                fields[self.start_field] = start
                fields["duedate"] = due
            if fields:
                if not self.dry_run:
                    self.client.jira_put(f"/issue/{existing.key}", {"fields": fields})
                self.report.updated_issues.append(f"{existing.key} (Epic {code}): {', '.join(fields)}")
        return out

    def push(self, data: PlanData, pair: dict[str, JiraIssueView], decisions: dict[str, dict[str, FieldDecision]],
             epic_by_code: dict[str, JiraIssueView], synced_ids: set[str]) -> dict[str, dict[str, str]]:
        rep = self.report
        synced = [t for t in data.tasks if t.sync]
        epic_keys = self.ensure_epics(data, epic_by_code, synced)
        epic_code_by_key = {k: c for c, k in epic_keys.items()}
        final_jira: dict[str, dict[str, str]] = {}
        sheet_url = self.backend.url
        project = self.jcfg["project_key"]

        for t in synced:
            v = pair.get(t.id)
            row_dec = decisions.get(t.id, {})
            owner_acc = self.resolve_owner(t.owner, data.people)
            group = data.groups.get(t.group)
            skip = {f for f, d in row_dec.items() if d.decision in ("conflict", "skip") or (d.decision == "pull" and not self.do_pull)}
            fields: dict[str, Any] = {}
            if v is None:
                fields = {
                    "project": {"key": project},
                    "issuetype": {"name": self.task_type},
                    "summary": f"{t.id}_{t.name}"[:255],
                    "description": build_issue_description(t, group, sheet_url),
                    "labels": self.desired_labels(t, []),
                }
                if t.group in epic_keys:
                    fields["parent"] = {"key": epic_keys[t.group]}
                if t.start:
                    fields[self.start_field] = t.start.isoformat()
                    fields["duedate"] = t.finish.isoformat()
                if owner_acc:
                    fields["assignee"] = {"accountId": owner_acc}
                if self.dry_run:
                    key = f"NEW-{t.id}"
                    rep.created_issues.append(f"{t.id} {t.name[:40]} (dry-run)")
                else:
                    created = self.client.jira_post("/issue", {"fields": fields})
                    key = created["key"]
                    rep.created_issues.append(f"{key} {t.id} {t.name[:40]}")
                    target = SHEET_TO_JIRA_STATUS.get(t.status, JIRA_TODO)
                    if target != JIRA_TODO:
                        self.transition(key, target)
                t.jira_key = key
                if not self.dry_run:
                    self.set_cell(t, "jira_key", key)
                pair[t.id] = JiraIssueView(key, fields["summary"], t.id, t.name,
                                           SHEET_TO_JIRA_STATUS.get(t.status, JIRA_TODO), fields["labels"],
                                           epic_keys.get(t.group, ""), owner_acc or "", "",
                                           t.start.isoformat() if t.start else "",
                                           t.finish.isoformat() if t.finish else "",
                                           norm_multiline(t.description), [], False)
                pair[t.id].desc_full = adf_to_text(fields["description"], stop_at_marker=False)
                continue

            changed = []
            want_summary = f"{t.id}_{t.name}"[:255]
            if "name" not in skip and v.summary != want_summary:
                fields["summary"] = want_summary
                changed.append("summary")
            if "description" not in skip:
                desc_adf = build_issue_description(t, group, sheet_url)
                if v.desc_full != adf_to_text(desc_adf, stop_at_marker=False):
                    fields["description"] = desc_adf
                    changed.append("description")
            labels = self.desired_labels(t, v.labels)
            if "priority" in skip:
                labels = [l for l in labels if not l.startswith("prio-")] + [l for l in v.labels if l.startswith("prio-")]
                labels = sorted(set(labels))
            if sorted(v.labels) != labels:
                fields["labels"] = labels
                changed.append("labels")
            want_parent = epic_keys.get(t.group, "")
            if "group" not in skip and want_parent and v.parent_key != want_parent:
                fields["parent"] = {"key": want_parent}
                changed.append("parent")
            want_start = t.start.isoformat() if t.start else ""
            want_due = t.finish.isoformat() if t.finish else ""
            if "start" not in skip and want_start and v.start != want_start:
                fields[self.start_field] = want_start
                changed.append("start")
            if "finish" not in skip and want_due and v.due != want_due:
                fields["duedate"] = want_due
                changed.append("due")
            dec_owner = row_dec.get("owner")
            if ("owner" not in skip and owner_acc and v.assignee_id != owner_acc
                    and (dec_owner is None or dec_owner.decision != "pull")):
                fields["assignee"] = {"accountId": owner_acc}
                changed.append("assignee")
            if fields:
                if not self.dry_run:
                    self.client.jira_put(f"/issue/{v.key}", {"fields": fields})
                rep.updated_issues.append(f"{v.key} {t.id}: {', '.join(changed)}")
            dec_status = row_dec.get("status")
            want_status = SHEET_TO_JIRA_STATUS.get(t.status, JIRA_TODO)
            if "status" not in skip and dec_status and dec_status.decision == "push" and v.status != want_status:
                if not self.dry_run:
                    self.transition(v.key, want_status)
                rep.updated_issues.append(f"{v.key} {t.id}: status → {want_status}")
                v.status = want_status
            # reflect pushed values into the view for snapshot computation
            if "summary" in fields:
                v.summary, v.name = want_summary, t.name
            if "description" in fields:
                v.desc_text = norm_multiline(t.description)
                v.desc_full = adf_to_text(fields["description"], stop_at_marker=False)
            if "labels" in fields:
                v.labels = labels
            if "parent" in fields:
                v.parent_key = want_parent
            if self.start_field in fields:
                v.start = want_start
            if "duedate" in fields:
                v.due = want_due
            if "assignee" in fields:
                v.assignee_id = owner_acc or ""

        # ----- dependency links -----
        id_to_key = {tid: vv.key for tid, vv in pair.items() if tid in synced_ids}
        key_to_id = {k: i for i, k in id_to_key.items()}
        for t in synced:
            v = pair.get(t.id)
            if v is None:
                continue
            row_dec = decisions.get(t.id, {})
            current = {key_to_id[l.other_key]: l for l in v.links if l.other_is_blocker and l.other_key in key_to_id}
            want = {p.id for p in t.preds if p.id in id_to_key}
            dec = row_dec.get("preds")
            snap_preds = set(filter(None, (dec.snap_value or "").split(","))) if dec and dec.snap_value is not None else None
            for pid in sorted(want - set(current)):
                self.add_link(id_to_key[pid], v.key, f"{pid} → {t.id}")
                v.links.append(JiraLink("new", id_to_key[pid], True))
            for pid in sorted(set(current) - want):
                # push-only：Jira 端新加、尚未拉回的相依不刪
                if not self.do_pull and snap_preds is not None and pid not in snap_preds:
                    continue
                link = current[pid]
                if not self.dry_run:
                    self.client.jira_delete(f"/issueLink/{link.link_id}")
                rep.links_removed.append(f"{pid} → {t.id} ({id_to_key[pid]} → {v.key})")
                v.links = [l for l in v.links if l is not link]

        key_to_id_all = {vv.key: tid for tid, vv in pair.items()}
        for t in synced:
            v = pair.get(t.id)
            if v is not None:
                final_jira[t.id] = self.jira_values(v, epic_code_by_key, key_to_id_all)
        return final_jira

    def transition(self, key: str, target: str) -> None:
        try:
            sr.transition_issue(self.client, key, target, dry_run=False)
        except Exception as exc:  # noqa: BLE001
            self.report.errors.append(f"{key} 狀態轉換到 {target} 失敗：{exc}")

    def add_link(self, pred_key: str, succ_key: str, label: str) -> None:
        rep = self.report
        if self.dry_run or pred_key.startswith("NEW") or succ_key.startswith("NEW"):
            rep.links_added.append(f"{label} (dry-run)" if self.dry_run else label)
            return
        payload = self._link_payload(pred_key, succ_key)
        self.client.jira_post("/issueLink", payload)
        if not self.link_verified:
            self.link_verified = True
            data = self.client.jira_get(f"/issue/{succ_key}", fields="issuelinks")
            ok = any(
                (l.get("type") or {}).get("name") == "Blocks" and (l.get("inwardIssue") or {}).get("key") == pred_key
                for l in (data.get("fields") or {}).get("issuelinks") or []
            )
            if not ok:
                for l in (data.get("fields") or {}).get("issuelinks") or []:
                    if (l.get("outwardIssue") or {}).get("key") == pred_key and (l.get("type") or {}).get("name") == "Blocks":
                        self.client.jira_delete(f"/issueLink/{l['id']}")
                self.link_flip = True
                self.client.jira_post("/issueLink", self._link_payload(pred_key, succ_key))
                rep.warnings.append("Jira issueLink 方向與預期相反，已自動修正")
        rep.links_added.append(f"{label} ({pred_key} blocks {succ_key})")

    def _link_payload(self, pred_key: str, succ_key: str) -> dict[str, Any]:
        inward, outward = (succ_key, pred_key) if self.link_flip else (pred_key, succ_key)
        return {"type": {"name": "Blocks"}, "inwardIssue": {"key": inward}, "outwardIssue": {"key": outward}}

    def handle_removed(self, candidates: list[tuple[str, JiraIssueView | None]], data: PlanData,
                       snap_items: dict[str, Any]) -> None:
        rep = self.report
        if not candidates:
            return
        limit = int(self.jcfg.get("max_removals") or 10)
        if len(candidates) > limit and not self.allow_mass_removal:
            rep.errors.append(
                f"本次有 {len(candidates)} 筆計畫單在 Sheet 中找不到（上限 {limit}），為安全起見未處理。"
                "確認無誤請加 --allow-mass-removal 重跑。"
            )
            return
        mode = self.jcfg.get("on_removed", "close")
        done = self.jcfg.get("done_status", JIRA_DONE)
        for key, v in candidates:
            desc = f"{key} {v.summary[:50] if v else ''}"
            if self.dry_run:
                (rep.closed_issues if mode == "close" else rep.deleted_issues).append(desc + " (dry-run)")
                continue
            try:
                if mode == "delete":
                    self.client.jira_delete(f"/issue/{key}?deleteSubtasks=false")
                    rep.deleted_issues.append(desc)
                else:
                    labels = sorted(set((v.labels if v else []) + ["plan-removed"]))
                    self.client.jira_put(f"/issue/{key}", {"fields": {"labels": labels}})
                    if v is None or v.status != done:
                        sr.transition_issue(self.client, key, done, dry_run=False)
                    rep.closed_issues.append(desc)
                snap_items.pop(key, None)
            except Exception as exc:  # noqa: BLE001
                rep.errors.append(f"處理已刪除的 {key} 失敗：{exc}")

    # ----- sheet write -----
    def write_sheet(self, data: PlanData) -> None:
        rep = self.report
        headers = list(data.task_headers)
        idx = header_index(headers)
        missing = [k for k in ("start", "finish", "jira_key", "conflict", "last_sync") if k not in idx]
        for k in missing:
            headers.append(TASK_COLUMNS[k][0])
        idx = header_index(headers)
        new_rows = [t for t in data.tasks if t.row == 0]
        updates = [(r, idx[c] + 1, v) for (r, c), v in sorted(self.sheet_updates.items()) if c in idx]
        user_cols = {idx[c] + 1 for c in DATE_COLUMNS | NUMERIC_COLUMNS | {"last_sync"} if c in idx}
        rep.sheet_writes = len(updates) + len(new_rows)
        if self.dry_run:
            if updates or new_rows:
                print(f"[dry-run] Sheet 將寫入 {len(updates)} 個儲存格、新增 {len(new_rows)} 列")
            return
        if not self.backend.writable:
            if updates or new_rows:
                rep.warnings.append(f"Sheet 唯讀：{len(updates)} 個儲存格與 {len(new_rows)} 列未寫回（請設定 service account）")
            return
        if missing:
            self.backend.ensure_headers(TASKS_TAB, headers)
        self.backend.write_cells(TASKS_TAB, updates, user_entered_cols=user_cols)
        if new_rows:
            rows = []
            for t in new_rows:
                row = [""] * len(headers)
                for k, i in idx.items():
                    row[i] = t.cells.get(k, "")
                row[idx["start"]] = t.start.isoformat() if t.start else ""
                row[idx["finish"]] = t.finish.isoformat() if t.finish else ""
                rows.append(row)
            self.backend.append_rows(TASKS_TAB, data.task_row_count + 1, rows)
        self.backend.write_settings({
            "Last Sync": datetime.now(TW_TZ).strftime("%Y-%m-%d %H:%M"),
            "Jira Project": f"{self.jcfg['project_key']} (label {self.label})",
            "Confluence Page": self.report.confluence_url or "(尚未建立)",
        })

    # ----- Confluence -----
    def update_confluence(self, data: PlanData, cal: WorkCalendar, state: dict[str, Any]) -> None:
        ccfg = self.pcfg["confluence"]
        body = render_confluence(data, cal, self.report, self.site_url, self.backend.url,
                                 show_discussion=bool(ccfg.get("show_discussion", True)))
        digest = hashlib.sha256(body.encode("utf-8")).hexdigest()
        stamp = datetime.now(TW_TZ).strftime("%Y-%m-%d %H:%M")
        full = body.replace("{{LAST_SYNC}}", html.escape(stamp))
        meta = state.setdefault("meta", {})
        title = str(ccfg.get("page_title") or "C27 Project Plan")
        page_id = str(ccfg.get("page_id") or meta.get("confluence_page_id") or "")
        if not page_id:
            page_id = self.find_child_page(str(ccfg["parent_page_id"]), title)
        if page_id in REGISTER_PAGE_IDS:
            raise PlanError("拒絕寫入 Register 頁面")
        space_key = ccfg.get("space_key", "projectmod")
        if page_id:
            self.report.confluence_url = f"{self.site_url}/wiki/spaces/{space_key}/pages/{page_id}"
            if meta.get("confluence_hash") == digest and meta.get("confluence_page_id") == page_id:
                self.report.confluence_action = "內容未變，略過"
                return
            if self.dry_run:
                self.report.confluence_action = "dry-run：將更新"
                return
            page = self.client.confluence_get(f"/pages/{page_id}")
            if page.get("title") in PROTECTED_PAGE_TITLES:
                raise PlanError("拒絕寫入 Register 頁面")
            self.client.confluence_put(f"/pages/{page_id}", {
                "id": page_id, "status": "current", "title": title,
                "body": {"representation": "storage", "value": full},
                "version": {"number": page["version"]["number"] + 1,
                            "message": ccfg.get("version_message", "Synced from C27 plan sheet"), "minorEdit": True},
            })
            self.report.confluence_action = "已更新（minorEdit）"
        else:
            if self.dry_run:
                self.report.confluence_action = "dry-run：將建立新頁面"
                return
            created = self.client.confluence_post("/pages", {
                "spaceId": str(ccfg.get("space_id") or "1409028"), "status": "current", "title": title,
                "parentId": str(ccfg["parent_page_id"]),
                "body": {"representation": "storage", "value": full},
            })
            page_id = str(created["id"])
            self.report.confluence_url = f"{self.site_url}/wiki/spaces/{space_key}/pages/{page_id}"
            self.report.confluence_action = "已建立"
        meta["confluence_page_id"] = page_id
        meta["confluence_hash"] = digest

    def find_child_page(self, parent_id: str, title: str) -> str:
        cursor = None
        for _ in range(20):
            params: dict[str, Any] = {"limit": 100}
            if cursor:
                params["cursor"] = cursor
            data = self.client.confluence_get(f"/pages/{parent_id}/children", **params)
            for p in data.get("results", []):
                if p.get("title") == title:
                    return str(p["id"])
            nxt = (data.get("_links") or {}).get("next", "")
            m = re.search(r"cursor=([^&]+)", nxt or "")
            if not m:
                break
            cursor = requests.utils.unquote(m.group(1))
        return ""


# ---------------------------------------------------------------------------
# Confluence rendering
# ---------------------------------------------------------------------------

def _status_macro(title: str, colour: str) -> str:
    return (
        '<ac:structured-macro ac:name="status"><ac:parameter ac:name="colour">'
        f"{colour}</ac:parameter><ac:parameter ac:name=\"title\">{html.escape(title)}</ac:parameter>"
        "</ac:structured-macro>"
    )


STATUS_COLOURS = {"Not Started": "Grey", "In Progress": "Blue", "Blocked": "Red", "Done": "Green", "Cancelled": "Grey"}


def _cell(text: str) -> str:
    return "<td>" + html.escape(text or "").replace("\n", "<br/>") + "</td>"


def render_confluence(data: PlanData, cal: WorkCalendar, rep: PlanReport, site_url: str, sheet_url: str,
                      *, show_discussion: bool = True) -> str:
    parts = []
    sheet_link = f'<a href="{html.escape(sheet_url)}">C27_Project_Plan（Google Sheet）</a>' if sheet_url else "(Google Sheet)"
    parts.append(
        '<ac:structured-macro ac:name="info"><ac:rich-text-body>'
        f"<p>計畫主檔：{sheet_link}　｜　最後同步：{{{{LAST_SYNC}}}} (UTC+8)</p>"
        "<p>本頁由 sync_plan.py 自動產生，請勿直接編輯；要修改計畫請改 Google Sheet。"
        "灰色「提案/討論中」列尚未建立 Jira 單（Sync to Jira = N）。工期單位：工作天（Mon–Fri，扣除台灣國定假日）。</p>"
        "</ac:rich-text-body></ac:structured-macro>"
    )
    parts.append("<h2>排程摘要</h2><ul>" + "".join(f"<li>{html.escape(l.strip())}</li>" for l in rep.summary_lines) + "</ul>")
    spans = group_spans(data.tasks)
    totals = {c: (e, a, ok) for c, e, a, ok in validate_group_totals(data.tasks, data.groups)}
    parts.append("<h2>群組（Epic）</h2><table><tbody><tr><th>Group</th><th>Epic</th><th>Start</th><th>Finish</th>"
                 "<th>任務工期合計 (d)</th><th>承諾工期</th></tr>")
    for code, g in data.groups.items():
        if code not in spans:
            continue
        s, f = spans[code]
        dur = sum(t.duration for t in data.tasks if t.group == code and t.type == TYPE_TASK)
        if code in totals:
            e, a, ok = totals[code]
            check = _status_macro(f"{e}d ✓" if ok else f"{e}d ≠ {a}d", "Green" if ok else "Red")
        else:
            check = _status_macro("估計", "Yellow")
        parts.append(f"<tr>{_cell(code)}{_cell(g.epic_name)}{_cell(str(s))}{_cell(str(f))}{_cell(str(dur))}<td>{check}</td></tr>")
    parts.append("</tbody></table>")
    if rep.conflicts or rep.warnings:
        items = [f"衝突：{c}" for c in rep.conflicts] + rep.warnings[:30]
        parts.append('<ac:structured-macro ac:name="warning"><ac:rich-text-body><ul>'
                     + "".join(f"<li>{html.escape(i)}</li>" for i in items) + "</ul></ac:rich-text-body></ac:structured-macro>")
    head = ["ID", "Group", "Task", "Type", "Priority", "Owner", "Dur (d)", "Predecessors", "Start", "Finish",
            "Status", "Sync", "Jira", "Description / 工作內容"]
    if show_discussion:
        head.append("Discussion")
    parts.append("<h2>任務清單</h2><table><tbody><tr>" + "".join(f"<th>{html.escape(h)}</th>" for h in head) + "</tr>")
    for t in data.tasks:
        jira = f'<a href="{site_url}/browse/{t.jira_key}">{t.jira_key}</a>' if t.jira_key and not t.jira_key.startswith("NEW") else ""
        sync = _status_macro("Jira", "Green") if t.sync else _status_macro("提案/討論中", "Grey")
        dur = f"{t.duration}" + (" (估)" if t.estimate else "")
        cells = [
            _cell(t.id), _cell(t.group), _cell(t.name), _cell(t.type), _cell(t.priority), _cell(t.owner or "(預設)"),
            _cell(dur), _cell(format_predecessors(t.preds)), _cell(t.start.isoformat() if t.start else ""),
            _cell(t.finish.isoformat() if t.finish else ""),
            f"<td>{_status_macro(t.status, STATUS_COLOURS.get(t.status, 'Grey'))}</td>", f"<td>{sync}</td>",
            f"<td>{jira}</td>", _cell(t.description),
        ]
        if show_discussion:
            cells.append(_cell(t.discussion))
        style = ' style="background-color: #f4f5f7;"' if not t.sync else ""
        parts.append(f"<tr{style}>" + "".join(cells) + "</tr>")
    parts.append("</tbody></table>")
    return "".join(parts)


# ---------------------------------------------------------------------------
# Sheet initialization / xlsx export
# ---------------------------------------------------------------------------

README_LINES = [
    ["C27 Project Plan — 使用說明"],
    [""],
    ["這份 Google Sheet 是 C27 計畫的主檔。大家在這裡討論（排程、要新增哪些工作、優先順序、工作內容），確認後按 run_sync_plan.bat 推到 Jira 與 Confluence。"],
    [""],
    ["【Tasks 欄位】"],
    ["ID：唯一且固定，格式 PLN-<群組>-<兩位數>，例 PLN-AT-01。新增列可留空，同步時自動配發。不要改既有 ID。"],
    ["Group：群組代碼（見 Groups 分頁），每個群組對應一張 Jira Epic。"],
    ["Type：Task / Milestone / Procurement。Duration = 0 視為 Milestone。"],
    ["Priority：P1 / P2 / P3，會變成 Jira label prio-p1 等（PMWC 沒有 priority 欄位）。"],
    ["Owner：People 分頁中的名字；空白 = 預設受託人 Shannon。"],
    ["Duration (d)：工作天（Mon–Fri，扣除 Holidays 分頁的假日）。1W = 5 天。"],
    ["Predecessors：Microsoft Project 語法，例 PLN-AT-01FS+1d; PLN-FX-02SS; PLN-WZ-03FF-2d。類型 FS/SS/FF/SF（省略 = FS），lag 以工作天 d 或週 w 表示，可為負，多個以 ; 或 , 分隔。"],
    ["Start Constraint：選填。填日期 = 不早於（SNET）；填「MSO 2026-11-02」= 必須在當天開始。"],
    ["Start / Finish：程式計算，請勿手改（要指定日期請填 Start Constraint）。"],
    ["Status：Not Started / In Progress / Blocked / Done / Cancelled（對應 Jira 待辦事項 / 進行中 / 進行中 / 完成 / 完成）。"],
    ["Sync to Jira：Y = 已核准，會建立/更新 Jira 單；N 或空白 = 提案/討論中，只顯示在 Confluence（灰色），不建 Jira 單。"],
    ["Jira Key：程式寫回，請勿修改。"],
    ["Description / 工作內容：詳細範圍，會寫進 Jira 描述。"],
    ["Discussion / Comments：討論用，不會進 Jira（會顯示在 Confluence 頁面）。"],
    ["Proposed By / Estimate / Notes：提案人、是否為估計值（Y = 待確認）、備註。只在 Sheet / Confluence。"],
    ["Sync Conflict：程式寫入衝突或資料錯誤。衝突時在此欄輸入 use sheet 或 use jira 決定以哪邊為準。"],
    [""],
    ["【雙向同步規則】"],
    ["1. 每次同步會記住上次同步時兩邊一致的值（隱藏分頁 _sync_state）。"],
    ["2. 只有一邊改 → 自動同步到另一邊；兩邊都改了同一欄 → 不覆寫任何一邊，寫在 Sync Conflict 欄並在郵件中列出。"],
    ["3. 計畫欄位（Task Name、Description、Duration/日期、Predecessors、Priority、Group、Sync flag）以 Sheet 為主；執行欄位（Status、Owner/受託人、Jira Key）以 Jira 為主。第一次比對沒有紀錄時以主方為準。"],
    ["4. 在 Jira timeline 拖曳日期 → 寫回 Sheet 的 Start Constraint（開始日）與 Duration（工期），再重新排程（仍會遵守相依關係）。"],
    ["5. 在 Jira 兩張計畫單之間新增/刪除 Blocks 相依 → 寫回 Predecessors（FS）。"],
    ["6. 在計畫 Epic 底下（或加 label c27-plan）新建的 Jira 單 → 自動加入 Sheet 新列並配發 ID。"],
    ["7. 從 Sheet 刪除整列 → Jira 單轉為「完成」並加 label plan-removed（不刪除，保留工時紀錄）。"],
    ["8. Sync to Jira 改回 N → Jira 單保留不動，也不再同步。"],
    [""],
    ["【執行】雙擊 run_sync_plan.bat（雙向）或 run_pull_plan_from_jira.bat（只把 Jira 變更拉回 Sheet）。"],
    ["GitHub Actions：bb00lin/my-bot → Sync Plan（sync-plan.yml），可選 direction = both / push / pull 與 dry_run。"],
]


def seed_tabs() -> dict[str, list[list[Any]]]:
    import plan_seed_c27 as seed

    tasks = [list(seed.TASK_HEADERS)]
    for r in seed.build_seed_tasks():
        tasks.append([r.get(h, "") for h in seed.TASK_HEADERS])
    groups = [["Code", "Epic Name (Jira)", "Locked Total (d)", "Description"]]
    groups += [[c, n, lt, d] for c, n, lt, d in seed.GROUPS]
    people = [["Name", "Email", "Jira Account ID", "Role"]] + [list(p) for p in seed.PEOPLE]
    holidays = [["Date", "Name", "Note"]] + [list(h) for h in seed.HOLIDAYS]
    settings = [
        ["Key", "Value", "說明"],
        ["Project Start", seed.PROJECT_START_PLACEHOLDER, "⚠ 佔位日期（下週一），請改成實際開工日。採購與 DOE/Carrier Board 從這天起算。"],
        ["Project Start Note", "placeholder", "設定實際日期後把這格改成 confirmed"],
        ["Working Days", "Mon,Tue,Wed,Thu,Fri", "工作日"],
        ["Jira Project", "PMWC (label c27-plan)", "（程式自動寫入，僅供顯示；實際設定在 GitHub secret PLAN_CONFIG_YAML）"],
        ["Confluence Parent", "C27_project_module_wifi7_QCC2072 (projectmod)", "（僅供顯示）"],
        ["Confluence Page", "(尚未建立)", "（程式自動寫入）"],
        ["On Removed", "close", "Sheet 刪除列時 Jira 單的處理方式（close / delete；實際設定在 PLAN_CONFIG_YAML）"],
        ["Last Sync", "", "（程式自動寫入）"],
    ]
    return {
        TASKS_TAB: tasks, GROUPS_TAB: groups, PEOPLE_TAB: people, HOLIDAYS_TAB: holidays,
        SETTINGS_TAB: settings, README_TAB: README_LINES,
    }


def compute_seed_schedule(tabs: dict[str, list[list[Any]]]) -> PlanData:
    data = build_plan_data(tabs)
    validate_ids(data.tasks)
    resolve_predecessors(data.tasks)
    compute_schedule(data.tasks, project_calendar(data))
    idx = header_index(data.task_headers)
    for t in data.tasks:
        row = tabs[TASKS_TAB][t.row - 1]
        row[idx["start"]] = t.start.isoformat() if t.start else ""
        row[idx["finish"]] = t.finish.isoformat() if t.finish else ""
    return data


GANTT_WEEKS = 40
GANTT_ROWS = 300


def gantt_values() -> list[list[Any]]:
    head = ["ID", "Task Name", "Start", "Finish"]
    first = "=Settings!B2-WEEKDAY(Settings!B2,3)"
    weeks = [first] + [f"={_col(5 + i - 1)}1+7" for i in range(1, GANTT_WEEKS)]
    rows = [head + weeks]
    for r in range(2, GANTT_ROWS + 1):
        rows.append([f'=IF(Tasks!A{r}="","",Tasks!A{r})', f'=IF(Tasks!A{r}="","",Tasks!D{r})',
                     f'=IF(Tasks!K{r}="","",Tasks!K{r})', f'=IF(Tasks!L{r}="","",Tasks!L{r})'])
    return rows


def _col(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def init_google_sheet(plan_cfg: dict[str, Any], *, force: bool = False) -> str:
    import gspread

    gc = gspread_client()
    if gc is None:
        raise PlanError("缺少 GOOGLE_SHEETS_JSON（service account），無法建立 Google Sheet")
    sheet_cfg = plan_cfg.get("sheet") or {}
    sid = sheet_id_from_url(str(sheet_cfg.get("sheet_id") or ""))
    title = str(sheet_cfg.get("title") or "C27_Project_Plan")
    sa = service_account_email()
    print(f"Service account: {sa}")
    sh = None
    try:
        sh = gc.open_by_key(sid) if sid else gc.open(title)
        print(f"找到既有試算表：{sh.title} ({sh.id})")
    except gspread.SpreadsheetNotFound:
        try:
            sh = gc.create(title)
            print(f"已建立試算表：{sh.id}")
        except gspread.exceptions.APIError as exc:
            raise PlanError(
                f"service account 無法建立檔案（{exc}）。請手動在 Google Drive 建立空白試算表「{title}」，"
                f"分享給 {sa}（編輯者），再重跑 init。"
            ) from exc
    titles = [w.title for w in sh.worksheets()]
    if TASKS_TAB in titles and not force:
        vals = sh.worksheet(TASKS_TAB).get_all_values()
        if len(vals) > 1:
            print("Tasks 分頁已有資料，略過初始化（要覆蓋請加 --force-init）")
            return sh.id
    tabs = seed_tabs()
    compute_seed_schedule(tabs)
    tabs[GANTT_TAB] = gantt_values()
    tabs[STATE_TAB] = [["key", "json"]]
    order = [TASKS_TAB, GROUPS_TAB, PEOPLE_TAB, HOLIDAYS_TAB, SETTINGS_TAB, GANTT_TAB, README_TAB, STATE_TAB]
    first = sh.sheet1
    ws_map = {}
    for i, name in enumerate(order):
        values = tabs[name]
        rows = max(len(values) + 50, 100)
        cols = max(max(len(r) for r in values) + 2, 6)
        if name in titles:
            ws = sh.worksheet(name)
            ws.clear()
            ws.resize(rows=rows, cols=cols)
        elif i == 0 and first.title not in order:
            ws = first
            ws.update_title(name)
            ws.resize(rows=rows, cols=cols)
        else:
            ws = sh.add_worksheet(name, rows=rows, cols=cols)
        ws.update(values=values, range_name="A1", value_input_option="USER_ENTERED")
        ws_map[name] = ws
    sh.batch_update({"requests": format_requests(ws_map, len(tabs[TASKS_TAB]))})
    for email in sheet_cfg.get("share_with") or []:
        try:
            sh.share(email, perm_type="user", role="writer", notify=True,
                     email_message="C27 Project Plan（計畫主檔）已建立，請直接在此討論與編輯。")
            print(f"已分享給 {email}（編輯者）")
        except Exception as exc:  # noqa: BLE001
            print(f"警告：分享給 {email} 失敗：{exc}")
    print(f"SHEET_URL=https://docs.google.com/spreadsheets/d/{sh.id}/edit")
    print(f"SHEET_ID={sh.id}")
    return sh.id


def format_requests(ws_map: dict[str, Any], task_rows: int) -> list[dict[str, Any]]:
    import plan_seed_c27 as seed

    tid = ws_map[TASKS_TAB].id
    hdr = seed.TASK_HEADERS
    col = {h: i for i, h in enumerate(hdr)}
    last_row = 1000
    reqs: list[dict[str, Any]] = [
        {"updateSpreadsheetProperties": {"properties": {"locale": "zh_TW", "timeZone": "Asia/Taipei"},
                                         "fields": "locale,timeZone"}},
        {"updateSheetProperties": {"properties": {"sheetId": tid, "gridProperties": {"frozenRowCount": 1, "frozenColumnCount": 4}},
                                   "fields": "gridProperties.frozenRowCount,gridProperties.frozenColumnCount"}},
        {"repeatCell": {"range": {"sheetId": tid, "startRowIndex": 0, "endRowIndex": 1},
                        "cell": {"userEnteredFormat": {"textFormat": {"bold": True, "foregroundColor": {"red": 1, "green": 1, "blue": 1}},
                                                       "backgroundColor": {"red": 0.2, "green": 0.33, "blue": 0.55},
                                                       "wrapStrategy": "WRAP", "verticalAlignment": "MIDDLE"}},
                        "fields": "userEnteredFormat(textFormat,backgroundColor,wrapStrategy,verticalAlignment)"}},
        {"setBasicFilter": {"filter": {"range": {"sheetId": tid, "startRowIndex": 0, "endRowIndex": last_row,
                                                 "startColumnIndex": 0, "endColumnIndex": len(hdr)}}}},
    ]
    widths = {"ID": 95, "WBS": 45, "Group": 55, "Task Name": 260, "Type": 95, "Priority": 60, "Owner": 110,
              "Duration (d)": 70, "Predecessors": 200, "Start Constraint": 100, "Start": 90, "Finish": 90,
              "Status": 100, "% Complete": 70, "Sync to Jira": 70, "Jira Key": 85,
              "Description / 工作內容": 320, "Discussion / Comments": 260, "Proposed By": 140, "Estimate": 65,
              "Notes": 220, "Sync Conflict": 220, "Last Sync": 120}
    for h, w in widths.items():
        reqs.append({"updateDimensionProperties": {"range": {"sheetId": tid, "dimension": "COLUMNS",
                                                             "startIndex": col[h], "endIndex": col[h] + 1},
                                                   "properties": {"pixelSize": w}, "fields": "pixelSize"}})

    def rng(h: str) -> dict[str, Any]:
        return {"sheetId": tid, "startRowIndex": 1, "endRowIndex": last_row, "startColumnIndex": col[h], "endColumnIndex": col[h] + 1}

    def one_of(h: str, values: list[str], strict: bool = True) -> dict[str, Any]:
        return {"setDataValidation": {"range": rng(h), "rule": {
            "condition": {"type": "ONE_OF_LIST", "values": [{"userEnteredValue": v} for v in values]},
            "strict": strict, "showCustomUi": True}}}

    def one_of_range(h: str, ref: str, strict: bool) -> dict[str, Any]:
        return {"setDataValidation": {"range": rng(h), "rule": {
            "condition": {"type": "ONE_OF_RANGE", "values": [{"userEnteredValue": ref}]},
            "strict": strict, "showCustomUi": True}}}

    reqs += [
        one_of("Type", [TYPE_TASK, TYPE_MILESTONE, TYPE_PROCUREMENT]),
        one_of("Priority", PRIORITY_OPTIONS),
        one_of("Status", STATUS_OPTIONS),
        one_of("Sync to Jira", ["Y", "N"]),
        one_of("Estimate", ["Y", "N"]),
        one_of_range("Group", f"='{GROUPS_TAB}'!$A$2:$A$100", True),
        one_of_range("Owner", f"='{PEOPLE_TAB}'!$A$2:$A$100", False),
    ]
    for h in ("Start Constraint", "Start", "Finish"):
        reqs.append({"repeatCell": {"range": rng(h), "cell": {"userEnteredFormat": {"numberFormat": {"type": "DATE", "pattern": "yyyy-mm-dd"}}},
                                    "fields": "userEnteredFormat.numberFormat"}})
    for h in ("Predecessors", "ID", "Jira Key"):
        reqs.append({"repeatCell": {"range": rng(h), "cell": {"userEnteredFormat": {"numberFormat": {"type": "TEXT"}}},
                                    "fields": "userEnteredFormat.numberFormat"}})
    for h in ("Task Name", "Description / 工作內容", "Discussion / Comments", "Notes", "Sync Conflict"):
        reqs.append({"repeatCell": {"range": rng(h), "cell": {"userEnteredFormat": {"wrapStrategy": "WRAP", "verticalAlignment": "TOP"}},
                                    "fields": "userEnteredFormat(wrapStrategy,verticalAlignment)"}})
    computed_bg = {"red": 0.93, "green": 0.95, "blue": 0.98}
    for h in ("Start", "Finish", "Jira Key", "Sync Conflict", "Last Sync"):
        reqs.append({"repeatCell": {"range": rng(h), "cell": {"userEnteredFormat": {"backgroundColor": computed_bg}},
                                    "fields": "userEnteredFormat.backgroundColor"}})
        reqs.append({"addProtectedRange": {"protectedRange": {"range": rng(h), "warningOnly": True,
                                                             "description": f"{h} 由程式寫入，手動修改會在下次同步被覆寫"}}})
    full = {"sheetId": tid, "startRowIndex": 1, "endRowIndex": last_row, "startColumnIndex": 0, "endColumnIndex": len(hdr)}
    sync_col = _col(col["Sync to Jira"] + 1)
    conf_col = _col(col["Sync Conflict"] + 1)
    type_col = _col(col["Type"] + 1)
    est_col = _col(col["Estimate"] + 1)
    reqs += [
        {"addConditionalFormatRule": {"index": 0, "rule": {"ranges": [full], "booleanRule": {
            "condition": {"type": "CUSTOM_FORMULA", "values": [{"userEnteredValue": f'=AND($A2<>"",${conf_col}2<>"")'}]},
            "format": {"backgroundColor": {"red": 1, "green": 0.85, "blue": 0.85}}}}}},
        {"addConditionalFormatRule": {"index": 1, "rule": {"ranges": [full], "booleanRule": {
            "condition": {"type": "CUSTOM_FORMULA", "values": [{"userEnteredValue": f'=AND($D2<>"",UPPER(${sync_col}2)<>"Y")'}]},
            "format": {"backgroundColor": {"red": 0.93, "green": 0.93, "blue": 0.93}, "textFormat": {"italic": True}}}}}},
        {"addConditionalFormatRule": {"index": 2, "rule": {"ranges": [full], "booleanRule": {
            "condition": {"type": "CUSTOM_FORMULA", "values": [{"userEnteredValue": f'=${type_col}2="Milestone"'}]},
            "format": {"textFormat": {"bold": True}}}}}},
        {"addConditionalFormatRule": {"index": 3, "rule": {"ranges": [rng("Duration (d)")], "booleanRule": {
            "condition": {"type": "CUSTOM_FORMULA", "values": [{"userEnteredValue": f'=${est_col}2="Y"'}]},
            "format": {"backgroundColor": {"red": 1, "green": 0.95, "blue": 0.7}}}}}},
    ]
    notes = {
        "ID": "唯一且固定：PLN-<群組>-<數字>。新增列可留空，同步時自動配發。",
        "Predecessors": "例：PLN-AT-01FS+1d; PLN-FX-02SS; PLN-WZ-03FF-2d（FS 可省略）",
        "Start Constraint": "選填。日期 = 不早於；MSO 2026-11-02 = 必須當天開始",
        "Start": "程式計算",
        "Finish": "程式計算",
        "Sync to Jira": "Y = 推到 Jira；N/空白 = 提案/討論中",
        "Discussion / Comments": "討論用，不進 Jira",
        "Sync Conflict": "衝突時輸入 use sheet 或 use jira",
    }
    for h, note in notes.items():
        reqs.append({"updateCells": {"range": {"sheetId": tid, "startRowIndex": 0, "endRowIndex": 1,
                                               "startColumnIndex": col[h], "endColumnIndex": col[h] + 1},
                                     "rows": [{"values": [{"note": note}]}], "fields": "note"}})
    for name in (GROUPS_TAB, PEOPLE_TAB, HOLIDAYS_TAB, SETTINGS_TAB):
        sid = ws_map[name].id
        reqs.append({"updateSheetProperties": {"properties": {"sheetId": sid, "gridProperties": {"frozenRowCount": 1}},
                                               "fields": "gridProperties.frozenRowCount"}})
        reqs.append({"repeatCell": {"range": {"sheetId": sid, "startRowIndex": 0, "endRowIndex": 1},
                                    "cell": {"userEnteredFormat": {"textFormat": {"bold": True}, "backgroundColor": {"red": 0.85, "green": 0.9, "blue": 0.97}}},
                                    "fields": "userEnteredFormat(textFormat,backgroundColor)"}})
        reqs.append({"updateDimensionProperties": {"range": {"sheetId": sid, "dimension": "COLUMNS", "startIndex": 0, "endIndex": 4},
                                                   "properties": {"pixelSize": 200}, "fields": "pixelSize"}})
    hid = ws_map[HOLIDAYS_TAB].id
    reqs.append({"repeatCell": {"range": {"sheetId": hid, "startRowIndex": 1, "endRowIndex": 500, "startColumnIndex": 0, "endColumnIndex": 1},
                                "cell": {"userEnteredFormat": {"numberFormat": {"type": "DATE", "pattern": "yyyy-mm-dd"}}},
                                "fields": "userEnteredFormat.numberFormat"}})
    stid = ws_map[SETTINGS_TAB].id
    reqs.append({"repeatCell": {"range": {"sheetId": stid, "startRowIndex": 1, "endRowIndex": 2, "startColumnIndex": 1, "endColumnIndex": 2},
                                "cell": {"userEnteredFormat": {"numberFormat": {"type": "DATE", "pattern": "yyyy-mm-dd"},
                                                               "backgroundColor": {"red": 1, "green": 0.9, "blue": 0.6}}},
                                "fields": "userEnteredFormat(numberFormat,backgroundColor)"}})
    gid = ws_map[GANTT_TAB].id
    reqs += [
        {"updateSheetProperties": {"properties": {"sheetId": gid, "gridProperties": {"frozenRowCount": 1, "frozenColumnCount": 4}},
                                   "fields": "gridProperties.frozenRowCount,gridProperties.frozenColumnCount"}},
        {"repeatCell": {"range": {"sheetId": gid, "startRowIndex": 0, "endRowIndex": 1, "startColumnIndex": 4, "endColumnIndex": 4 + GANTT_WEEKS},
                        "cell": {"userEnteredFormat": {"numberFormat": {"type": "DATE", "pattern": "mm/dd"}, "textRotation": {"angle": 90}}},
                        "fields": "userEnteredFormat(numberFormat,textRotation)"}},
        {"repeatCell": {"range": {"sheetId": gid, "startRowIndex": 1, "endRowIndex": GANTT_ROWS, "startColumnIndex": 2, "endColumnIndex": 4},
                        "cell": {"userEnteredFormat": {"numberFormat": {"type": "DATE", "pattern": "yyyy-mm-dd"}}},
                        "fields": "userEnteredFormat.numberFormat"}},
        {"updateDimensionProperties": {"range": {"sheetId": gid, "dimension": "COLUMNS", "startIndex": 4, "endIndex": 4 + GANTT_WEEKS},
                                       "properties": {"pixelSize": 24}, "fields": "pixelSize"}},
        {"updateDimensionProperties": {"range": {"sheetId": gid, "dimension": "COLUMNS", "startIndex": 1, "endIndex": 2},
                                       "properties": {"pixelSize": 240}, "fields": "pixelSize"}},
        {"addConditionalFormatRule": {"index": 0, "rule": {"ranges": [{"sheetId": gid, "startRowIndex": 1, "endRowIndex": GANTT_ROWS,
                                                                        "startColumnIndex": 4, "endColumnIndex": 4 + GANTT_WEEKS}],
                                                           "booleanRule": {"condition": {"type": "CUSTOM_FORMULA", "values": [
                                                               {"userEnteredValue": "=AND(ISNUMBER($C2),E$1<=$D2,E$1+6>=$C2)"}]},
                                                               "format": {"backgroundColor": {"red": 0.27, "green": 0.51, "blue": 0.71}}}}}},
        {"updateSheetProperties": {"properties": {"sheetId": ws_map[STATE_TAB].id, "hidden": True}, "fields": "hidden"}},
        {"updateDimensionProperties": {"range": {"sheetId": ws_map[README_TAB].id, "dimension": "COLUMNS", "startIndex": 0, "endIndex": 1},
                                       "properties": {"pixelSize": 1100}, "fields": "pixelSize"}},
        {"repeatCell": {"range": {"sheetId": ws_map[README_TAB].id, "startColumnIndex": 0, "endColumnIndex": 1},
                        "cell": {"userEnteredFormat": {"wrapStrategy": "WRAP"}}, "fields": "userEnteredFormat.wrapStrategy"}},
    ]
    return reqs


def export_xlsx(path: Path) -> PlanData:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation

    import plan_seed_c27 as seed

    tabs = seed_tabs()
    data = compute_seed_schedule(tabs)
    tabs[GANTT_TAB] = gantt_values()
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="335A8C")
    for name in (TASKS_TAB, GROUPS_TAB, PEOPLE_TAB, HOLIDAYS_TAB, SETTINGS_TAB, GANTT_TAB, README_TAB):
        ws = wb.create_sheet(name)
        for row in tabs[name]:
            out = []
            for v in row:
                d = parse_date_any(v) if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v or "") else None
                if d:
                    out.append(d)
                elif isinstance(v, str) and re.fullmatch(r"-?\d+", v):
                    out.append(int(v))
                else:
                    out.append(v)
            ws.append(out)
        if name != README_TAB:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = header_fill
                c.alignment = Alignment(wrap_text=True, vertical="center")
    ws = wb[TASKS_TAB]
    ws.freeze_panes = "E2"
    hdr = seed.TASK_HEADERS
    col = {h: _col(i + 1) for i, h in enumerate(hdr)}
    widths = {"ID": 13, "Task Name": 40, "Predecessors": 30, "Description / 工作內容": 50, "Discussion / Comments": 35,
              "Notes": 30, "Sync Conflict": 30, "Proposed By": 20, "Owner": 14, "Type": 12, "Status": 12}
    for h, c in col.items():
        ws.column_dimensions[c].width = widths.get(h, 11)
    for h in ("Start Constraint", "Start", "Finish"):
        for cell in ws[col[h]][1:]:
            cell.number_format = "yyyy-mm-dd"
    for h in ("Task Name", "Description / 工作內容", "Discussion / Comments", "Notes"):
        for cell in ws[col[h]][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    lists = {
        "Type": [TYPE_TASK, TYPE_MILESTONE, TYPE_PROCUREMENT], "Priority": PRIORITY_OPTIONS,
        "Status": STATUS_OPTIONS, "Sync to Jira": ["Y", "N"], "Estimate": ["Y", "N"],
    }
    for h, opts in lists.items():
        dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col[h]}2:{col[h]}1000")
    dv = DataValidation(type="list", formula1=f"='{GROUPS_TAB}'!$A$2:$A$100", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col['Group']}2:{col['Group']}1000")
    dv = DataValidation(type="list", formula1=f"='{PEOPLE_TAB}'!$A$2:$A$100", allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(f"{col['Owner']}2:{col['Owner']}1000")
    ws.conditional_formatting.add(f"A2:{_col(len(hdr))}1000", FormulaRule(
        formula=[f'AND($D2<>"",UPPER(${col["Sync to Jira"]}2)<>"Y")'], fill=PatternFill("solid", fgColor="EEEEEE")))
    ws.conditional_formatting.add(f"{col['Duration (d)']}2:{col['Duration (d)']}1000", FormulaRule(
        formula=[f'${col["Estimate"]}2="Y"'], fill=PatternFill("solid", fgColor="FFF2B3")))
    ws.auto_filter.ref = f"A1:{_col(len(hdr))}{len(tabs[TASKS_TAB])}"
    g = wb[GANTT_TAB]
    g.freeze_panes = "E2"
    g.column_dimensions["B"].width = 40
    for i in range(GANTT_WEEKS):
        g.column_dimensions[_col(5 + i)].width = 4
        g.cell(row=1, column=5 + i).number_format = "mm/dd"
    for r in range(2, GANTT_ROWS + 1):
        g.cell(row=r, column=3).number_format = "yyyy-mm-dd"
        g.cell(row=r, column=4).number_format = "yyyy-mm-dd"
    g.conditional_formatting.add(f"E2:{_col(4 + GANTT_WEEKS)}{GANTT_ROWS}", FormulaRule(
        formula=["AND(ISNUMBER($C2),E$1<=$D2,E$1+6>=$C2)"], fill=PatternFill("solid", fgColor="4682B4")))
    wb[README_TAB].column_dimensions["A"].width = 160
    for name in (GROUPS_TAB, PEOPLE_TAB, HOLIDAYS_TAB, SETTINGS_TAB):
        wb[name].freeze_panes = "A2"
        for c in "ABCD":
            wb[name].column_dimensions[c].width = 30
    wb.save(path)
    return data


# ---------------------------------------------------------------------------
# Email
# ---------------------------------------------------------------------------

def send_report_email(cfg: dict[str, Any], rep: PlanReport, body: str) -> None:
    notify = cfg.get("notify") or {}
    if rep.dry_run or not sr.is_github_actions() or not notify.get("email_on_github", True):
        return
    verdict = "衝突" if rep.conflicts else ("有變更" if rep.has_changes() else "無變更")
    if rep.errors:
        verdict = "錯誤"
    subject = f"[C27 Plan Sync] {verdict} {datetime.now(TW_TZ):%Y-%m-%d %H:%M}"
    html_body = "<pre style=\"font-family:Consolas,monospace;font-size:13px\">" + html.escape(body) + "</pre>"
    try:
        sr.send_diff_email(cfg, subject=subject, body=body, html_body=html_body)
        print(f"結果郵件已寄出：{subject}")
    except Exception as exc:  # noqa: BLE001
        print(f"警告：寄送結果郵件失敗：{exc}")
        rep.warnings.append(f"寄送結果郵件失敗：{exc}")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def make_client(cfg: dict[str, Any]) -> PlanAtlassianClient:
    a = cfg["atlassian"]
    return PlanAtlassianClient(a["site_url"], a["cloud_id"], a["email"], cfg["_api_token"],
                               silent_mode=bool(cfg.get("silent_mode", True)))


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="C27 Project Plan：Google Sheet <-> Jira PMWC -> Confluence")
    parser.add_argument("--config", type=Path, default=SCRIPT_DIR / "config.yaml")
    parser.add_argument("--plan-config", type=Path, default=SCRIPT_DIR / "plan_config.yaml")
    parser.add_argument("--direction", choices=["both", "push", "pull"], default="both")
    parser.add_argument("--dry-run", action="store_true", help="只預覽，不寫入 Sheet / Jira / Confluence")
    parser.add_argument("--xlsx", default="", help="改從本機 xlsx（或公開匯出 URL）唯讀讀取")
    parser.add_argument("--init-sheet", action="store_true", help="建立/初始化 Google Sheet（需 GOOGLE_SHEETS_JSON）")
    parser.add_argument("--force-init", action="store_true", help="初始化時覆蓋既有 Tasks 分頁")
    parser.add_argument("--export-xlsx", type=Path, help="輸出初始內容的 .xlsx 範本後結束")
    parser.add_argument("--allow-mass-removal", action="store_true")
    parser.add_argument("--report-file", type=Path, default=SCRIPT_DIR / "logs" / "plan_sync_latest.txt")
    args = parser.parse_args(argv)

    try:
        if args.export_xlsx:
            data = export_xlsx(args.export_xlsx)
            print(f"已輸出 {args.export_xlsx}（{len(data.tasks)} 筆任務）")
            return 0
        if args.init_sheet:
            cfg = load_plan_config(args.config, args.plan_config, need_token=False)
            init_google_sheet(cfg["plan"], force=args.force_init)
            return 0
        cfg = load_plan_config(args.config, args.plan_config)
        backend = open_backend(cfg["plan"], args.xlsx)
        client = make_client(cfg)
        sync = PlanSync(cfg, backend, client, direction=args.direction, dry_run=args.dry_run,
                        allow_mass_removal=args.allow_mass_removal)
        rep = sync.run()
    except PlanError as exc:
        print(f"同步中止：{exc}", file=sys.stderr)
        return 2
    body = rep.render()
    print(body)
    args.report_file.parent.mkdir(parents=True, exist_ok=True)
    args.report_file.write_text(body, encoding="utf-8")
    send_report_email(cfg, rep, body)
    return 1 if rep.errors else 0


if __name__ == "__main__":
    sys.exit(main())
